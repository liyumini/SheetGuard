"""gold_v2:完整错误检测测试集的载体工作簿。

在现有 build_gold_model(seed=0) 的 5 张表上追加(原有行列一律不动):

- Rate 表:VLOOKUP 查找表(A 列键 1..12,B 列折算系数常量);
- RawData 行 6:MonthNo 常量 1..12(VLOOKUP 查找键);
- Revenue 行 3:IF 横向家族(条件比较 + 数值分支);
- Revenue 行 4:VLOOKUP 横向家族(装饰行,不作为载体——绝对范围
  Rate!A2:B13 使横向模板互不相同,见探针边界 2);
- Agg 表:5 行明细(横向家族,各引用不同源表) + N 列行尾聚合
  MIN/MAX/AVERAGE/SUMIF/COUNTIF(聚合格豁免判据生效,跨表聚合必
  误报 dependency_anomaly,故聚合只引用本表,见探针边界 1);
- Lk 表:B 列纵向 VLOOKUP 家族(键 C{r}/表 C{r}:D{r} 随行平移,
  模板一致,是 VLOOKUP 的唯一故障载体,见探针边界 2)。

生成时守门(full_gate 断言在 complete_cases.py 生成入口):
① 全表重算零异常;② 干净静态候选集为空。
"""
from __future__ import annotations

import os
from tempfile import NamedTemporaryFile

from openpyxl import Workbook

from evaluation.generation.workbook import build_gold_model
from sheetguard.spreadsheet.anomaly_detector import StaticInspector
from sheetguard.spreadsheet.dependency_graph import DependencyGraph
from sheetguard.spreadsheet.parser import parse_workbook
from sheetguard.spreadsheet.recalc import RecalcEngine

SUBSET_FUNCTIONS = ("SUM", "IF", "VLOOKUP", "SUMIF", "COUNTIF", "AVERAGE", "MIN", "MAX")

# 与 build_gold_model 的月份布局一致:n_months=12,CS=2(B),Total 列 = N
_COLS = ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]


def build_gold_v2_model(seed: int = 0) -> Workbook:
    """在原 gold 上追加子集函数载体行;原行列零改动。"""
    wb = build_gold_model(seed=seed)

    # ── Rate 查找表 ──
    rate = wb.create_sheet("Rate")
    rate["A1"] = "MonthNo"
    rate["B1"] = "Factor"
    for i in range(1, 13):
        rate.cell(i + 1, 1, i)
        rate.cell(i + 1, 2, round(1.0 - (i - 1) * 0.01, 2))

    # ── RawData 行 6:MonthNo 常量(查找键)──
    raw = wb["RawData"]
    raw["A6"] = "MonthNo"
    for idx, col in enumerate(_COLS):
        raw[f"{col}6"] = idx + 1

    # ── Revenue 行 3-4:IF / VLOOKUP 横向家族 ──
    rev = wb["Revenue"]
    for col in _COLS:
        rev[f"{col}3"] = (
            f"=IF(RawData!{col}2>1000, Revenue!{col}2*0.9, Revenue!{col}2)"
        )
        rev[f"{col}4"] = (
            f"=VLOOKUP(RawData!{col}6, Rate!A2:B13, 2, FALSE)"
        )
    rev["A3"] = "DiscountedGross"
    rev["A4"] = "SeasonFactor"
    for r in (3, 4):
        rev[f"N{r}"] = f"=SUM(B{r}:M{r})"

    # ── Agg 表:5 行明细 + N 列行尾聚合(全部本表引用)──
    agg = wb.create_sheet("Agg")
    agg["A1"] = "Label"
    for m in range(1, 13):
        agg.cell(1, 1 + m, f"M{m}")
    agg.cell(1, 14, "Total")
    detail = [
        ("Revenue", 2), ("Cost", 2), ("P&L", 2),
        ("Dashboard", 2), ("RawData", 2),
    ]
    for row_idx, (src, src_row) in enumerate(detail, start=2):
        agg.cell(row_idx, 1, f"Line{row_idx - 1}")
        for col in _COLS:
            agg[f"{col}{row_idx}"] = f"='{src}'!{col}{src_row}"
    for r, formula in {
        2: "=MIN(B2:M2)",
        3: "=MAX(B3:M3)",
        4: "=AVERAGE(B4:M4)",
        # 条件聚合载体必须"值可反推"：SUMIF/COUNTIF 的常量不能同时被
        # 无条件聚合（SUM/COUNT）复现，否则 formula_accuracy 退化成抽奖
        # （09-15 complete 实验 complete_009/010 的教训，见
        # complete_cases._assert_gold_repairable 的唯一性守门）。
        5: '=SUMIF(B5:M5, ">=700", B5:M5)',
        6: '=COUNTIF(B6:M6, ">1000")',
    }.items():
        agg[f"N{r}"] = formula

    # ── Lk 表:B 列纵向 VLOOKUP 家族(VLOOKUP 载体)──
    lk = wb.create_sheet("Lk")
    lk["A1"] = "FactorLookup"
    for r in range(2, 14):
        lk.cell(r, 3, r - 1)                       # C 列键常量 1..12
        lk.cell(r, 4, round(0.5 + (r - 1) * 0.01, 2))  # D 列系数常量
        lk[f"B{r}"] = f"=VLOOKUP(C{r}, C{r}:D{r}, 2, FALSE)"
    return wb


def _gate_index(wb) -> tuple:
    """保存到临时文件并解析出 (index, graph, engine),供守门函数复用。"""
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        tmp_path = fh.name
    wb.save(tmp_path)
    try:
        index = parse_workbook(tmp_path)
        graph = DependencyGraph(index)
        graph.build()
        engine = RecalcEngine(index, graph)
        return index, graph, engine
    finally:
        os.unlink(tmp_path)


def recalc_failures(wb) -> list[tuple[str, str]]:
    """对全表公式逐格重算,返回 (完整地址, 异常描述) 列表;空列表 = 通过。"""
    index, _graph, engine = _gate_index(wb)
    failures: list[tuple[str, str]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    addr = f"{ws.title}!{cell.coordinate}"
                    try:
                        engine.evaluate_cell(addr)
                    except Exception as exc:
                        failures.append((addr, f"{type(exc).__name__}: {exc}"))
    return failures


def clean_candidates(wb) -> list[str]:
    """干净工作簿的静态候选集;空列表 = 通过(对照 hard_003)。"""
    index, _graph, _engine = _gate_index(wb)
    return sorted(c["cell"] for c in StaticInspector(index).detect_all())
