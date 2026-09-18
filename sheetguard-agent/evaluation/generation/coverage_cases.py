"""coverage 数据集:公式/表达式形态 × 5 类错误类型的覆盖矩阵。

与 complete_cases 同构(errors/expected_status/static_candidate_count +
gold_path/broken_path),全部 expected_status=success,生成期走同一套守门:
静态候选集必须与注入目标精确一致 + gold 必须通过 Verifier 六项验证
(missing_formula 另加常量唯一性守门)。

覆盖矩阵(9 个重算引擎函数 + 四则表达式 × 5 类错误,共 33 个单对案例):
- SUM / AVERAGE / MIN / MAX   wrong_range + missing_formula + cross_sheet_error
- COUNT                        wrong_range + missing_formula(跨表对值中性,见下)
- IF                           wrong_operator(比较翻转) / wrong_cell_reference /
                               missing_formula / cross_sheet_error
- VLOOKUP                      wrong_range(查表收缩) / wrong_cell_reference(键位移) /
                               missing_formula(Lk B) / cross_sheet_error(Lk F)
- SUMIF                        wrong_range / wrong_cell_reference(sum_range 错列) /
                               missing_formula / cross_sheet_error
- COUNTIF                      wrong_range / missing_formula / cross_sheet_error
- Arith(四则表达式)            wrong_operator(+/- 互换) / wrong_cell_reference /
                               missing_formula / cross_sheet_error
cov_033 为五故障混合(SUM×wrong_range + Arith×wrong_operator + IF×
missing_formula + VLOOKUP×wrong_cell_reference + COUNTIF×cross_sheet_error),
一条工作簿内 5 类错误 × 5 种载体各占其一。

标 N/A 的组合(矩阵如实留空,不凑数):
- wrong_range × IF:IF 无范围引用;
- wrong_cell_reference × 纯范围聚合(SUM/AVERAGE/MIN/MAX/COUNT):只有范围
  引用,"引用错位"语义并入 wrong_range;SUMIF 的 sum_range 错列可表达该对;
- wrong_operator × SUMIF/COUNTIF:criteria 运算符在字符串常量里,5 个静态
  信号结构上不可见,行尾聚合格豁免又兜住模板偏差——静态不可锚定(探针实测,
  complete-v1 的 SUMIF/COUNTIF 载体也只覆盖 missing_formula);
- wrong_operator × 纯聚合:公式内无运算符;
- cross_sheet_error × COUNT:对任意全数值行计数恒等于格数(值中性,调查阶段
  会被按"值正确"dismiss,语义上不构成修复问题);
- cross_sheet_error × VLOOKUP 本地家族(Lk B):无跨表引用,由新增 Lk F
  跨表家族覆盖。

表达式形态随载体一并覆盖:比较运算符(IF 条件)、字符串常量(SUMIF/COUNTIF
criteria)、布尔常量(VLOOKUP 第 4 参 FALSE)。绝对引用 $ 沿用 complete-v1
探针边界(绝对范围不作故障载体)。
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from evaluation.generation.complete_cases import (
    _assert_candidates,
    _assert_gold_repairable,
    _hardcode,
)
from evaluation.generation.coverage_workbook import build_gold_v3_model

# 手写注入规格:(case_id, faults);faults = (sheet, coord, broken_formula,
# error_type, carrier_function, expression_form)
_HAND_SPECS: list[tuple[str, list[tuple[str, str, str, str, str, str]]]] = [
    ("cov_000", [  # SUM × wrong_range
        ("Revenue", "N3", "=SUM(B3:L3)", "wrong_range", "SUM", "range_aggregate"),
    ]),
    ("cov_002", [  # SUM × cross_sheet_error
        ("Revenue", "N3", "=SUM(Cost!B3:M3)", "cross_sheet_error", "SUM", "range_aggregate"),
    ]),
    ("cov_003", [  # AVERAGE × wrong_range
        ("Agg", "N4", "=AVERAGE(B4:L4)", "wrong_range", "AVERAGE", "range_aggregate"),
    ]),
    ("cov_005", [  # AVERAGE × cross_sheet_error
        ("Agg", "N4", "=AVERAGE(Revenue!B4:M4)", "cross_sheet_error", "AVERAGE",
         "range_aggregate"),
    ]),
    ("cov_006", [  # MIN × wrong_range
        ("Agg", "N2", "=MIN(B2:L2)", "wrong_range", "MIN", "range_aggregate"),
    ]),
    ("cov_008", [  # MIN × cross_sheet_error
        ("Agg", "N2", "=MIN(Cost!B2:M2)", "cross_sheet_error", "MIN", "range_aggregate"),
    ]),
    ("cov_009", [  # MAX × wrong_range
        ("Agg", "N3", "=MAX(B3:L3)", "wrong_range", "MAX", "range_aggregate"),
    ]),
    ("cov_011", [  # MAX × cross_sheet_error
        ("Agg", "N3", "=MAX(Revenue!B3:M3)", "cross_sheet_error", "MAX", "range_aggregate"),
    ]),
    ("cov_012", [  # COUNT × wrong_range
        ("Agg", "N7", "=COUNT(B7:L7)", "wrong_range", "COUNT", "range_aggregate"),
    ]),
    ("cov_014", [  # IF × wrong_operator(比较翻转)
        ("Revenue", "D3", "=IF(RawData!D2<1000, Revenue!D2*0.9, Revenue!D2)",
         "wrong_operator", "IF", "conditional"),
    ]),
    ("cov_015", [  # IF × wrong_cell_reference(条件引用错位)
        ("Revenue", "D3", "=IF(RawData!E2>1000, Revenue!D2*0.9, Revenue!D2)",
         "wrong_cell_reference", "IF", "conditional"),
    ]),
    ("cov_017", [  # IF × cross_sheet_error
        ("Revenue", "D3", "=IF(Dashboard!D2>1000, Revenue!D2*0.9, Revenue!D2)",
         "cross_sheet_error", "IF", "conditional"),
    ]),
    ("cov_018", [  # VLOOKUP × wrong_range(查表收缩 → col_index 越界)
        ("Lk", "B7", "=VLOOKUP(C7, C7:C7, 2, FALSE)", "wrong_range", "VLOOKUP", "lookup"),
    ]),
    ("cov_019", [  # VLOOKUP × wrong_cell_reference(查找键位移)
        ("Lk", "B7", "=VLOOKUP(C8, C7:D7, 2, FALSE)", "wrong_cell_reference",
         "VLOOKUP", "lookup"),
    ]),
    ("cov_021", [  # VLOOKUP × cross_sheet_error(Lk F 跨表家族,Rate→RawData)
        ("Lk", "F7", "=VLOOKUP(C7, RawData!A7:B7, 2, FALSE)", "cross_sheet_error",
         "VLOOKUP", "lookup"),
    ]),
    ("cov_022", [  # SUMIF × wrong_range(criteria 与 sum_range 同步收缩)
        ("Agg", "N5", '=SUMIF(B5:L5, ">=700", B5:L5)', "wrong_range", "SUMIF",
         "conditional_aggregate"),
    ]),
    ("cov_023", [  # SUMIF × wrong_cell_reference(sum_range 错列一列)
        ("Agg", "N5", '=SUMIF(B5:M5, ">=700", C5:M5)', "wrong_cell_reference",
         "SUMIF", "conditional_aggregate"),
    ]),
    ("cov_025", [  # SUMIF × cross_sheet_error
        ("Agg", "N5", '=SUMIF(Revenue!B5:M5, ">=700", Revenue!B5:M5)',
         "cross_sheet_error", "SUMIF", "conditional_aggregate"),
    ]),
    ("cov_026", [  # COUNTIF × wrong_range
        ("Agg", "N6", '=COUNTIF(B6:L6, ">1000")', "wrong_range", "COUNTIF",
         "conditional_aggregate"),
    ]),
    ("cov_028", [  # COUNTIF × cross_sheet_error
        ("Agg", "N6", '=COUNTIF(Revenue!B6:M6, ">1000")', "cross_sheet_error",
         "COUNTIF", "conditional_aggregate"),
    ]),
    ("cov_029", [  # Arith × wrong_operator(减号换加号)
        ("P&L", "C2", "=Revenue!C2+Cost!C2", "wrong_operator", "Arith", "arithmetic"),
    ]),
    ("cov_030", [  # Arith × wrong_cell_reference(引用错位一格)
        ("Revenue", "E2", "=RawData!F2*RawData!E3", "wrong_cell_reference",
         "Arith", "arithmetic"),
    ]),
    ("cov_032", [  # Arith × cross_sheet_error(Cost 行引用整体换源)
        ("Cost", "E2", "=Revenue!E2*Revenue!E4+Revenue!E5", "cross_sheet_error",
         "Arith", "arithmetic"),
    ]),
]

# 硬编码规格:(case_id, coord, carrier_function, expression_form)
_HC_SPECS: list[tuple[str, str, str, str]] = [
    ("cov_001", "Revenue!N3", "SUM", "range_aggregate"),
    ("cov_004", "Agg!N4", "AVERAGE", "range_aggregate"),
    ("cov_007", "Agg!N2", "MIN", "range_aggregate"),
    ("cov_010", "Agg!N3", "MAX", "range_aggregate"),
    ("cov_013", "Agg!N7", "COUNT", "range_aggregate"),
    ("cov_016", "Revenue!D3", "IF", "conditional"),
    ("cov_020", "Lk!B7", "VLOOKUP", "lookup"),
    ("cov_024", "Agg!N5", "SUMIF", "conditional_aggregate"),
    ("cov_027", "Agg!N6", "COUNTIF", "conditional_aggregate"),
    ("cov_031", "P&L!C2", "Arith", "arithmetic"),
]

# cov_033 五故障混合:5 类错误 × 5 种载体,单一工作簿内同时注入。
_MIX_SPEC = {
    "case_id": "cov_033",
    "hardcodes": [("Revenue!D3", "IF", "conditional")],
    "faults": [
        ("Revenue", "N3", "=SUM(B3:L3)", "wrong_range", "SUM", "range_aggregate"),
        ("P&L", "C2", "=Revenue!C2+Cost!C2", "wrong_operator", "Arith", "arithmetic"),
        ("Lk", "B9", "=VLOOKUP(C10, C9:D9, 2, FALSE)", "wrong_cell_reference",
         "VLOOKUP", "lookup"),
        ("Agg", "N6", '=COUNTIF(Revenue!B6:M6, ">1000")', "cross_sheet_error",
         "COUNTIF", "conditional_aggregate"),
    ],
}


def _gold_formulas(gold_path: Path, targets: list[tuple[str, str]]) -> dict[str, str]:
    """读取 gold 工作簿中指定单元格的原始公式,作为标准答案。"""
    wb = load_workbook(gold_path)
    return {f"{sheet}!{coord}": str(wb[sheet][coord].value) for sheet, coord in targets}


def _build_spec_case(gold_path: Path, broken_dir: Path, spec: dict) -> dict:
    """按规格生成单案例:手写注入 + 硬编码,两道守门全部通过才返回记录。"""
    broken_path = broken_dir / f"{spec['case_id']}.xlsx"
    shutil.copy2(gold_path, broken_path)
    wb = load_workbook(broken_path)

    old_values: dict[str, str] = {}
    for coord, fn, form in spec.get("hardcodes", []):
        value = _hardcode(str(gold_path), coord)
        ws, col = coord.split("!")
        old_values[coord] = str(value)
        wb[ws][col] = value
    for sheet, coord, broken, *_ in spec.get("faults", []):
        wb[sheet][coord] = broken
    wb.save(broken_path)

    targets = {c for c, *_ in spec.get("hardcodes", [])} | {
        f"{s}!{c}" for s, c, *_ in spec.get("faults", [])
    }
    _assert_candidates(broken_path, targets)

    gold_formulas = _gold_formulas(
        gold_path, [tuple(t.split("!")) for t in sorted(targets)]
    )
    errors = [
        {
            "target_cell": coord,
            "error_type": "missing_formula",
            "old_formula": old_values[coord],
            "gold_formula": gold_formulas[coord],
            "carrier_function": fn,
            "expression_form": form,
        }
        for coord, fn, form in spec.get("hardcodes", [])
    ] + [
        {
            "target_cell": f"{sheet}!{coord}",
            "error_type": error_type,
            "old_formula": broken,
            "gold_formula": gold_formulas[f"{sheet}!{coord}"],
            "carrier_function": fn,
            "expression_form": form,
        }
        for sheet, coord, broken, error_type, fn, form in spec.get("faults", [])
    ]
    _assert_gold_repairable(broken_path, errors)

    return {
        "case_id": spec["case_id"],
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": errors,
        "expected_status": "success",
        "static_candidate_count": len(errors),
    }


def build_coverage_cases(gold_path: str, broken_dir: str, output_json: str) -> list[dict]:
    """构造覆盖矩阵数据集并写出清单 JSON;每个案例先过候选/可解性双守门。"""
    gold_path = Path(gold_path)
    broken_dir = Path(broken_dir)
    broken_dir.mkdir(parents=True, exist_ok=True)
    if not gold_path.exists():
        build_gold_v3_model(seed=0).save(gold_path)

    specs: list[dict] = [
        {"case_id": case_id, "faults": faults} for case_id, faults in _HAND_SPECS
    ] + [
        {
            "case_id": case_id,
            "hardcodes": [(coord, fn, form)],
        }
        for case_id, coord, fn, form in _HC_SPECS
    ] + [_MIX_SPEC]

    specs.sort(key=lambda s: s["case_id"])
    # 唯一性守门:同 case_id 只允许一个规格(手写/硬编码混排时防重复键)。
    ids = [s["case_id"] for s in specs]
    if len(set(ids)) != len(ids):
        raise AssertionError(f"case_id 重复: {sorted({i for i in ids if ids.count(i) > 1})}")

    records = [_build_spec_case(gold_path, broken_dir, spec) for spec in specs]

    output = Path(output_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    return records


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[2]
    gold = root / "evaluation/data/coverage/gold.xlsx"
    cases = build_coverage_cases(
        str(gold),
        str(root / "evaluation/data/coverage/broken"),
        str(root / "evaluation/data/coverage_cases.json"),
    )
    print(f"coverage cases: {len(cases)}")
    for case in cases:
        carriers = sorted({e["carrier_function"] for e in case["errors"]})
        etypes = sorted({e["error_type"] for e in case["errors"]})
        print(
            f"  {case['case_id']}: errors={len(case['errors'])} "
            f"carriers={'+'.join(carriers)} etypes={'+'.join(etypes)}"
        )
