"""complete 数据集:21 案例矩阵(6 信号专属 + 8 函数载体 + 组合)。

案例结构与 hard_cases 一致(errors/expected_status/static_candidate_count,
另附 gold_path/broken_path);每个案例生成期守门:静态候选集必须与注入
目标精确一致,保证 gold 期望可判定。全部 expected_status=success。

案例矩阵(与设计文档一致,全部探针验证):
- complete_000..005  信号专属:硬编码 / 引用错位 / 邻居不一致 / 范围收窄 /
  孤立汇总泄漏 / 跨表换源,各锚定一种静态信号;
- complete_006..013  函数载体:with_function(inject_missing_formula, fn)
  × 8 个公式子集函数(SUM/IF/VLOOKUP/SUMIF/COUNTIF/AVERAGE/MIN/MAX);
- complete_014..017  交叉载体:比较符翻转 / 键位移 / 跨表换源 / 范围收缩;
- complete_018       五类注入器混合(种子搜索,循环排除);
- complete_019       双断链(Revenue/Cost 各断一处);
- complete_020       三函数载体格混合硬编码(IF+VLOOKUP+MIN)。
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from evaluation.generation.complete_injectors import (
    has_function,
    inject_neighbor_mismatch,
    inject_singleton_sum_leak,
    with_function,
)
from evaluation.generation.complete_workbook import SUBSET_FUNCTIONS, build_gold_v2_model
from evaluation.generation.fault_injection import (
    inject_cross_sheet_error,
    inject_faults,
    inject_missing_formula,
    inject_wrong_cell_reference,
    inject_wrong_operator,
    inject_wrong_range,
)
from sheetguard.spreadsheet.anomaly_detector import StaticInspector
from sheetguard.spreadsheet.dependency_graph import DependencyGraph
from sheetguard.spreadsheet.parser import parse_workbook
from sheetguard.spreadsheet.recalc import RecalcEngine
from sheetguard.spreadsheet.value_hints import aggregate_value_hints
from sheetguard.spreadsheet.verifier import Verifier


def _gold_formulas(gold_path: Path, targets: list[tuple[str, str]]) -> dict[str, str]:
    """读取 gold 工作簿中指定单元格的原始公式,作为标准答案。"""
    wb = load_workbook(gold_path)
    return {
        f"{sheet}!{coord}": str(wb[sheet][coord].value)
        for sheet, coord in targets
    }


def _error(gold_formulas: dict[str, str], target: str, error_type: str, old_formula: str) -> dict:
    """组装一条错误级 gold 记录(old_formula = broken 副本写入后的内容)。"""
    return {
        "target_cell": target,
        "error_type": error_type,
        "old_formula": old_formula,
        "gold_formula": gold_formulas[target],
    }


def _assert_candidates(path: Path, expected: set[str]) -> None:
    """生成期守门:静态候选集必须与设计一致,否则该案例不可信。"""
    idx = parse_workbook(str(path))
    got = {row["cell"] for row in StaticInspector(idx).detect_all()}
    if got != expected:
        raise AssertionError(
            f"{path.name}: 静态候选集与设计不一致 got={sorted(got)} expected={sorted(expected)}"
        )


def _assert_gold_repairable(broken_path: Path, errors: list[dict]) -> None:
    """gold 可解性 + 唯一性守门（P1-4）。

    ① 可解性：逐条把 gold 公式写回 broken 副本，必须通过 Verifier 全部
       六项检查——gold 过不了的案例对 Agent 无解（complete_018 的
       Revenue!L4：无 $ 锚定查找区让 gold 被伪主流模板结构性判死）。
    ② 唯一性：missing_formula 的原常量必须由 gold 公式唯一复现——
       无条件聚合 gold（MIN/MAX/...）之外还有别的聚合同值，或条件聚合
       gold（SUMIF/COUNTIF）被某个无条件聚合复现（complete_009 的
       SUMIF==SUM），formula_accuracy 都会退化成抽奖。

    两个守门与 value_hints/Verifier 共用同一判据，保证"生成期认为可解"
    与"Agent 运行时认为可解"口径一致。
    """
    src = parse_workbook(str(broken_path))
    graph = DependencyGraph(src)
    graph.build()
    verifier = Verifier()
    for err in errors:
        target, gold = err["target_cell"], err["gold_formula"]
        attempt_path = _patched_temp_copy(broken_path, target, gold)
        try:
            result = verifier.verify(src, str(attempt_path), target, working_index=src)
        finally:
            Path(attempt_path).unlink(missing_ok=True)
        if not result.passed:
            raise AssertionError(
                f"{broken_path.name}: {target} 的 gold 公式未通过验证 "
                f"{result.checks}（gold 不可解：验证器规则与数据设计矛盾）"
            )
        if err["error_type"] != "missing_formula":
            continue
        hints = aggregate_value_hints(src, graph, target)
        matches = [h["formula"] for h in hints if h["matches_constant"]]
        gold_is_no_criteria = gold.replace(" ", "") in {h["formula"] for h in hints}
        if gold_is_no_criteria:
            if len(matches) != 1:
                raise AssertionError(
                    f"{broken_path.name}: {target} 的 gold 值可被多个聚合复现 "
                    f"{matches}，gold 不唯一"
                )
        elif matches:
            raise AssertionError(
                f"{broken_path.name}: {target} 的 gold 是条件聚合，但 "
                f"{matches} 也能复现原常量，gold 不唯一"
            )


def _patched_temp_copy(broken_path: Path, target: str, formula: str) -> str:
    """把 formula 写入 broken 副本的指定单元格，返回临时文件路径。"""
    import os
    import tempfile

    fd, temp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb = load_workbook(str(broken_path))
    sheet, coord = target.split("!")
    wb[sheet][coord] = formula
    wb.save(temp_path)
    return temp_path


def _hardcode(gold_path: Path, target: str) -> float:
    """用 gold 的 RecalcEngine 求指定单元格的正确数值(同 missing_formula 注入)。"""
    idx = parse_workbook(str(gold_path))
    graph = DependencyGraph(idx)
    graph.build()
    engine = RecalcEngine(idx, graph)
    return engine.evaluate_cell(target)


def _handwritten_case(gold_path: Path, broken_dir: Path, case_id: str,
                      faults: list[tuple[str, str, str, str]]) -> dict:
    """手写公式注入:faults = [(sheet, coord, broken_formula, error_type)]。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    shutil.copy2(gold_path, broken_path)
    wb = load_workbook(broken_path)
    for sheet, coord, broken_formula, _error_type in faults:
        wb[sheet][coord] = broken_formula
    wb.save(broken_path)
    _assert_candidates(broken_path, {f"{s}!{c}" for s, c, *_ in faults})
    gold_formulas = _gold_formulas(gold_path, [(s, c) for s, c, *_ in faults])
    _assert_gold_repairable(broken_path, [
        _error(gold_formulas, f"{s}!{c}", et, bf) for (s, c, bf, et) in faults
    ])
    return {
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, f"{s}!{c}", et, bf)
            for (s, c, bf, et) in faults
        ],
        "expected_status": "success",
        "static_candidate_count": len(faults),
    }


def _hardcode_case(gold_path: Path, broken_dir: Path, case_id: str, coords: list[str]) -> dict:
    """硬编码注入:把 gold 的 RecalcEngine 重算值写入 broken 副本。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    shutil.copy2(gold_path, broken_path)
    wb = load_workbook(broken_path)
    old_values: dict[str, str] = {}
    for coord in coords:
        value = _hardcode(gold_path, coord)
        ws, c = coord.split("!")
        old_values[coord] = str(value)
        wb[ws][c] = value
    wb.save(broken_path)
    _assert_candidates(broken_path, set(coords))
    gold_formulas = _gold_formulas(
        gold_path, [tuple(coord.split("!")) for coord in coords],
    )
    _assert_gold_repairable(broken_path, [
        _error(gold_formulas, coord, "missing_formula", old_values[coord])
        for coord in coords
    ])
    return {
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, coord, "missing_formula", old_values[coord])
            for coord in coords
        ],
        "expected_status": "success",
        "static_candidate_count": len(coords),
    }


def _injector_case(gold_path: Path, broken_dir: Path, case_id: str,
                   injector, max_seeds: int = 30) -> dict:
    """注入器案例:种子自检直到守门通过。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    for seed in range(0, max_seeds):
        try:
            meta = injector(str(gold_path), str(broken_path), seed=seed, copy_gold=True)
        except ValueError:
            continue
        try:
            _assert_candidates(broken_path, {meta.target_cell})
        except AssertionError:
            continue
        try:
            _assert_gold_repairable(broken_path, [{
                "target_cell": meta.target_cell,
                "error_type": meta.error_type,
                "old_formula": meta.old_formula,
                "gold_formula": meta.gold_formula,
            }])
        except AssertionError:
            continue
        return {
            "case_id": case_id,
            "gold_path": str(gold_path),
            "broken_path": str(broken_path),
            "errors": [
                {
                    "target_cell": meta.target_cell,
                    "error_type": meta.error_type,
                    "old_formula": meta.old_formula,
                    "gold_formula": meta.gold_formula,
                }
            ],
            "expected_status": "success",
            "static_candidate_count": 1,
        }
    raise RuntimeError(f"{case_id}: {max_seeds} 个种子内守门不通过")


def _carrier_case(gold_path: Path, broken_dir: Path, case_id: str, fn: str) -> dict:
    """函数载体:with_function(inject_missing_formula, fn) + 守门。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    for seed in range(0, 40):
        try:
            meta = with_function(inject_missing_formula, fn)(
                str(gold_path), str(broken_path), seed=seed, copy_gold=True,
            )
        except ValueError:
            continue
        if not has_function(meta.gold_formula, fn):
            continue
        try:
            _assert_candidates(broken_path, {meta.target_cell})
        except AssertionError:
            continue
        try:
            _assert_gold_repairable(broken_path, [{
                "target_cell": meta.target_cell,
                "error_type": meta.error_type,
                "old_formula": meta.old_formula,
                "gold_formula": meta.gold_formula,
            }])
        except AssertionError:
            continue
        return {
            "case_id": case_id,
            "gold_path": str(gold_path),
            "broken_path": str(broken_path),
            "errors": [
                {
                    "target_cell": meta.target_cell,
                    "error_type": meta.error_type,
                    "old_formula": meta.old_formula,
                    "gold_formula": meta.gold_formula,
                }
            ],
            "expected_status": "success",
            "static_candidate_count": 1,
        }
    raise RuntimeError(f"{case_id}: {fn} 载体在 40 个种子内守门不通过")


def _five_fault_case(gold_path: Path, broken_dir: Path, case_id: str = "complete_018") -> dict:
    """五类混合:inject_faults 五注入器 + 循环排除 + 候选守门(hard_004 模式)。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    specs_base = [
        (inject_wrong_range, 11),
        (inject_wrong_operator, 7),
        (inject_missing_formula, 81),
        (inject_wrong_cell_reference, 42),
    ]
    injected = None
    for cross_sheet_seed in range(1, 200):
        specs = [*specs_base, (inject_cross_sheet_error, cross_sheet_seed)]
        try:
            candidates = inject_faults(str(gold_path), str(broken_path), specs)
        except RuntimeError:
            continue
        idx = parse_workbook(str(broken_path))
        graph = DependencyGraph(idx)
        graph.build()
        if graph.cycle_sccs():
            continue
        try:
            _assert_candidates(broken_path, {m.target_cell for m in candidates})
        except AssertionError:
            continue
        try:
            _assert_gold_repairable(broken_path, [
                {
                    "target_cell": m.target_cell,
                    "error_type": m.error_type,
                    "old_formula": m.old_formula,
                    "gold_formula": m.gold_formula,
                }
                for m in candidates
            ])
        except AssertionError:
            continue
        injected = candidates
        break
    if injected is None:
        raise RuntimeError(f"{case_id}: 200 个种子内找不到无循环且 gold 可解的组合")
    return {
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            {
                "target_cell": m.target_cell,
                "error_type": m.error_type,
                "old_formula": m.old_formula,
                "gold_formula": m.gold_formula,
            }
            for m in injected
        ],
        "expected_status": "success",
        "static_candidate_count": len(injected),
    }


def _double_chain_case(gold_path: Path, broken_dir: Path, case_id: str = "complete_019") -> dict:
    """双断链:Revenue!F2 与 Cost!F2 各断一处,汇入 P&L!F2。"""
    broken_path = broken_dir / f"{case_id}.xlsx"
    shutil.copy2(gold_path, broken_path)
    wb = load_workbook(broken_path)
    wb["Revenue"]["F2"] = "=RawData!F2+RawData!F3"
    wb["Cost"]["F2"] = "=RawData!F2*RawData!F4-RawData!F5"
    wb.save(broken_path)
    _assert_candidates(broken_path, {"Revenue!F2", "Cost!F2"})
    gold_formulas = _gold_formulas(gold_path, [("Revenue", "F2"), ("Cost", "F2")])
    _assert_gold_repairable(broken_path, [
        _error(gold_formulas, "Revenue!F2", "wrong_operator",
               "=RawData!F2+RawData!F3"),
        _error(gold_formulas, "Cost!F2", "wrong_operator",
               "=RawData!F2*RawData!F4-RawData!F5"),
    ])
    return {
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, "Revenue!F2", "wrong_operator",
                   "=RawData!F2+RawData!F3"),
            _error(gold_formulas, "Cost!F2", "wrong_operator",
                   "=RawData!F2*RawData!F4-RawData!F5"),
        ],
        "expected_status": "success",
        "static_candidate_count": 2,
    }


def _three_fn_case(gold_path: Path, broken_dir: Path, case_id: str = "complete_020") -> dict:
    """三函数混合硬编码:IF/VLOOKUP/MIN 载体格写 gold 重算值。"""
    coords = ("Revenue!E3", "Lk!B7", "Agg!N2")
    broken_path = broken_dir / f"{case_id}.xlsx"
    shutil.copy2(gold_path, broken_path)
    gold_wb = load_workbook(gold_path)
    wb = load_workbook(broken_path)
    old_values: dict[str, str] = {}
    for coord in coords:
        value = _hardcode(gold_path, coord)
        ws, c = coord.split("!")
        old_values[coord] = str(value)
        wb[ws][c] = value
    wb.save(broken_path)
    _assert_candidates(broken_path, set(coords))
    gold_formulas = {
        f"{ws}!{c}": str(gold_wb[ws][c].value)
        for ws, c in (coord.split("!") for coord in coords)
    }
    _assert_gold_repairable(broken_path, [
        {"target_cell": t, "error_type": "missing_formula",
         "old_formula": old_values[t], "gold_formula": gold_formulas[t]}
        for t in coords
    ])
    return {
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            {"target_cell": t, "error_type": "missing_formula",
             "old_formula": old_values[t], "gold_formula": gold_formulas[t]}
            for t in coords
        ],
        "expected_status": "success",
        "static_candidate_count": len(coords),
    }


def build_complete_cases(gold_path: str, broken_dir: str, output_json: str) -> list[dict]:
    """构造 21 个完整错误检测案例并写出清单 JSON;每个案例先做候选守门。"""
    gold_path = Path(gold_path)
    broken_dir = Path(broken_dir)
    broken_dir.mkdir(parents=True, exist_ok=True)
    if not gold_path.exists():
        build_gold_v2_model(seed=0).save(gold_path)

    records: list[dict] = []

    # ── complete_000..005 信号专属 ──
    records.append(_hardcode_case(gold_path, broken_dir, "complete_000", ["Revenue!H2"]))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_001", [
        ("Revenue", "F2", "=RawData!G2*RawData!F3", "wrong_cell_reference"),
    ]))
    records.append(_injector_case(gold_path, broken_dir, "complete_002", inject_neighbor_mismatch))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_003", [
        ("Revenue", "N3", "=SUM(B3:L3)", "wrong_range"),
    ]))
    records.append(_injector_case(gold_path, broken_dir, "complete_004", inject_singleton_sum_leak))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_005", [
        ("P&L", "G2", "=Revenue!G2-RawData!G2", "cross_sheet_error"),
    ]))

    # ── complete_006..013 函数载体(with_function × 8)──
    for i, fn in enumerate(SUBSET_FUNCTIONS, start=6):
        records.append(_carrier_case(gold_path, broken_dir, f"complete_{i:03d}", fn))

    # ── complete_014..017 交叉载体 ──
    records.append(_handwritten_case(gold_path, broken_dir, "complete_014", [
        ("Revenue", "E3", "=IF(RawData!E2<1000, Revenue!E2*0.9, Revenue!E2)", "wrong_operator"),
    ]))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_015", [
        ("Lk", "B7", "=VLOOKUP(C8, C7:D7, 2, FALSE)", "wrong_cell_reference"),
    ]))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_016", [
        ("Agg", "E5", "=Revenue!E2", "cross_sheet_error"),
    ]))
    records.append(_handwritten_case(gold_path, broken_dir, "complete_017", [
        ("Agg", "N4", "=AVERAGE(B4:L4)", "wrong_range"),
    ]))

    # ── complete_018 五故障混合(种子搜索,含循环排除)──
    records.append(_five_fault_case(gold_path, broken_dir))

    # ── complete_019 双断链 ──
    records.append(_double_chain_case(gold_path, broken_dir))

    # ── complete_020 三函数混合硬编码 ──
    records.append(_three_fn_case(gold_path, broken_dir))

    output = Path(output_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    return records


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[2]
    gold = root / "evaluation/data/complete/gold.xlsx"
    cases = build_complete_cases(
        str(gold),
        str(root / "evaluation/data/complete/broken"),
        str(root / "evaluation/data/complete_cases.json"),
    )
    print(f"complete cases: {len(cases)}")
    for case in cases:
        print(
            f"  {case['case_id']}: errors={len(case['errors'])} "
            f"status={case['expected_status']} candidates={case['static_candidate_count']}"
        )
