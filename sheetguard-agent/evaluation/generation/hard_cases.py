"""multi-cell 难题数据集生成（补充 v1.6 高阶能力的判别性案例）。

现有 multi-cell 数据集全部是 "expected_status=success、2-3 个可修复错误"
的规整案例，deferred / dependency-skip / dismissal / false-repair 等评测
器从未被真正检验过（所有案例 issues=[]、各计数全零）。本模块构造六类
难题，每个案例先用 StaticInspector 离线验证候选集与信号，保证 gold
期望（该修什么、该跳过什么、该搁置什么、该拒绝什么）可判定：

- ``hard_000`` 误报陷阱：一个交换操作数的正确公式（结构异常、值正确）
  + 一个真实错误；期望 Agent 调查后 dismiss 陷阱、修复真错误。
- ``hard_001`` 依赖循环：Dashboard 两格互相引用形成 cycle（调度器必须
  skipped_due_to_dependency）+ 两个 missing_formula 真错误。
- ``hard_002`` 预算延期：12 个同分 missing_formula 硬编码，超过
  max_candidates=10 的部分必须 deferred_due_to_budget（同分并列按地址
  排序决定 admit 集合，完全确定）。
- ``hard_003`` 干净工作簿：零候选，Agent 不得幻觉修复（0 次修复）。
- ``hard_004`` 五类错误混合：五类注入器各注入一个错误，一次全修对。
- ``hard_005`` 双断链：Revenue/Cost 各断一处，共同汇入 P&L!F2，考察
  多上游修复后下游的级联验证。

所有 case 共享 ``build_gold_model`` 生成的 gold 工作簿；broken 副本写入
``evaluation/data/broken_hard/``，清单写入 ``evaluation/data/hard_cases.json``
（结构同 multi_cases.json，另含可选 expected_dismissed / expected_skipped /
expected_deferred 三个 workbook 级期望键）。
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from evaluation.generation.fault_injection import (
    inject_cross_sheet_error,
    inject_faults,
    inject_missing_formula,
    inject_wrong_cell_reference,
    inject_wrong_operator,
    inject_wrong_range,
)
from evaluation.generation.workbook import build_gold_model
from sheetguard.spreadsheet.anomaly_detector import StaticInspector
from sheetguard.spreadsheet.dependency_graph import DependencyGraph
from sheetguard.spreadsheet.parser import parse_workbook
from sheetguard.spreadsheet.recalc import RecalcEngine

# 与 graph.DEFAULT_MAX_CANDIDATES 保持一致；hard_002 依赖该预算值。
_MAX_CANDIDATES = 10

_INJECTORS = {
    "wrong_range": inject_wrong_range,
    "wrong_cell_reference": inject_wrong_cell_reference,
    "wrong_operator": inject_wrong_operator,
    "missing_formula": inject_missing_formula,
    "cross_sheet_error": inject_cross_sheet_error,
}


def _gold_formulas(gold_path: Path, targets: list[tuple[str, str]]) -> dict[str, str]:
    """读取 gold 工作簿中指定单元格的原始公式，作为标准答案。"""
    wb = load_workbook(gold_path)
    return {
        f"{sheet}!{coord}": str(wb[sheet][coord].value)
        for sheet, coord in targets
    }


def _error(
    gold_formulas: dict[str, str],
    target: str,
    error_type: str,
    old_formula: str,
) -> dict:
    """组装一条错误级 gold 记录。"""
    return {
        "target_cell": target,
        "error_type": error_type,
        "old_formula": old_formula,
        "gold_formula": gold_formulas[target],
    }


def _assert_candidates(path: Path, expected: set[str]) -> None:
    """生成期守门：静态候选集必须与设计一致，否则该案例不可信。"""
    index = parse_workbook(str(path))
    got = {row["cell"] for row in StaticInspector(index).detect_all()}
    if got != expected:
        raise AssertionError(
            f"{path.name}: 静态候选集与设计不一致 got={sorted(got)} expected={sorted(expected)}"
        )


def _copy_gold(gold_path: Path, broken_path: Path) -> None:
    shutil.copy2(gold_path, broken_path)


def _hardcoded_value(gold_path: Path, target: str) -> float:
    """用 gold 的 RecalcEngine 求指定单元格的正确数值（同 missing_formula 注入）。"""
    index = parse_workbook(str(gold_path))
    graph = DependencyGraph(index)
    graph.build()
    engine = RecalcEngine(index, graph)
    return engine.evaluate_cell(target)


def build_hard_cases(gold_path: str, broken_dir: str, output_json: str) -> list[dict]:
    """构造六个难题案例并写出清单 JSON；每个案例先做候选集守门。"""
    gold_path = Path(gold_path)
    broken_dir = Path(broken_dir)
    broken_dir.mkdir(parents=True, exist_ok=True)
    if not gold_path.exists():
        build_gold_model(seed=0).save(gold_path)

    records: list[dict] = []

    # ── hard_000 误报陷阱 ────────────────────────────────────────────
    case_id = "hard_000"
    broken_path = broken_dir / f"{case_id}.xlsx"
    _copy_gold(gold_path, broken_path)
    wb = load_workbook(broken_path)
    # 陷阱：交换操作数，结构异常但值正确；Agent 应在调查后 dismiss。
    wb["Revenue"]["E2"] = "=RawData!E3*RawData!E2"
    # 真错误：P&L!C2 的减号换成加号。
    wb["P&L"]["C2"] = "=Revenue!C2+Cost!C2"
    wb.save(broken_path)
    _assert_candidates(broken_path, {"P&L!C2", "Revenue!E2"})
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(
                _gold_formulas(gold_path, [("P&L", "C2")]),
                "P&L!C2", "wrong_operator", "=Revenue!C2+Cost!C2",
            ),
        ],
        "expected_status": "success",
        "expected_dismissed": ["Revenue!E2"],
        "static_candidate_count": 2,
    })

    # ── hard_001 依赖循环 ────────────────────────────────────────────
    case_id = "hard_001"
    broken_path = broken_dir / f"{case_id}.xlsx"
    _copy_gold(gold_path, broken_path)
    wb = load_workbook(broken_path)
    # 互相引用形成 cycle；调度器必须把两格 skipped_due_to_dependency。
    wb["Dashboard"]["F3"] = "=Dashboard!F4+1"
    wb["Dashboard"]["F4"] = "=Dashboard!F3-1"
    # 硬编码值必须取 gold 的正确计算值：value 检查以"原值"为基准，
    # 任意假值会让正确修复公式重算必然对不上，案例变成无解。
    wb["P&L"]["E2"] = _hardcoded_value(gold_path, "P&L!E2")
    wb["P&L"]["H2"] = _hardcoded_value(gold_path, "P&L!H2")
    wb.save(broken_path)
    _assert_candidates(
        broken_path, {"Dashboard!F3", "Dashboard!F4", "P&L!E2", "P&L!H2"},
    )
    gold_formulas = _gold_formulas(
        gold_path, [("P&L", "E2"), ("P&L", "H2")],
    )
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, "P&L!E2", "missing_formula",
                   str(_hardcoded_value(gold_path, "P&L!E2"))),
            _error(gold_formulas, "P&L!H2", "missing_formula",
                   str(_hardcoded_value(gold_path, "P&L!H2"))),
        ],
        "expected_status": "partial_success",
        "expected_skipped": ["Dashboard!F3", "Dashboard!F4"],
        "static_candidate_count": 4,
    })

    # ── hard_002 预算延期 ────────────────────────────────────────────
    case_id = "hard_002"
    broken_path = broken_dir / f"{case_id}.xlsx"
    _copy_gold(gold_path, broken_path)
    targets = [
        ("Revenue", "D2"), ("Revenue", "G2"), ("Revenue", "J2"),
        ("Cost", "D2"), ("Cost", "G2"), ("Cost", "K2"),
        ("P&L", "D2"), ("P&L", "G2"), ("P&L", "J2"), ("P&L", "L2"),
        ("Dashboard", "D2"), ("Dashboard", "G2"),
    ]
    gold_formulas = _gold_formulas(gold_path, targets)
    # 12 个候选同分（missing_formula=0.6），并列时按地址升序 admit 前
    # _MAX_CANDIDATES 个，其余 deferred_due_to_budget —— 与调度器的
    # (-score, -confidence, address) 排序完全一致。
    ordered = sorted(f"{sheet}!{coord}" for sheet, coord in targets)
    deferred = ordered[_MAX_CANDIDATES:]
    wb = load_workbook(broken_path)
    # 用 gold 的 RecalcEngine 求正确值再硬编码（同 missing_formula 注入）。
    index = parse_workbook(str(gold_path))
    graph = DependencyGraph(index)
    graph.build()
    engine = RecalcEngine(index, graph)
    hardcoded: dict[str, float] = {}
    for sheet, coord in targets:
        address = f"{sheet}!{coord}"
        hardcoded[address] = engine.evaluate_cell(address)
        wb[sheet][coord] = hardcoded[address]
    wb.save(broken_path)
    _assert_candidates(
        broken_path, {f"{sheet}!{coord}" for sheet, coord in targets},
    )
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, address, "missing_formula", str(value))
            for address, value in hardcoded.items()
        ],
        "expected_status": "partial_success",
        # 预算延期 = 规划期 2 个（地址排序尾部）+ 处理期 3 个链式
        # budget_dependency：被延期的 Revenue!G2/J2 是 P&L!G2/J2 提案的
        # 引用上游，P&L!G2 被延期后又链住 Dashboard!G2（设计 8.2 不提交
        # 上游未定的修复）。该链由规划决策与 precommit 规则确定，完全可复现。
        "expected_deferred": sorted({
            *deferred,
            "P&L!G2", "P&L!J2", "Dashboard!G2",
        }),
        "static_candidate_count": len(targets),
    })

    # ── hard_003 干净工作簿 ──────────────────────────────────────────
    case_id = "hard_003"
    broken_path = broken_dir / f"{case_id}.xlsx"
    _copy_gold(gold_path, broken_path)
    _assert_candidates(broken_path, set())
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [],
        # 零种子时 inspect 直接落 "no_candidates"（graph.py inspect 节点），
        # 不经过 plan_batch 的 "no_eligible_candidates" 分支。
        "expected_status": "no_candidates",
        "static_candidate_count": 0,
    })

    # ── hard_004 五类错误混合 ────────────────────────────────────────
    case_id = "hard_004"
    broken_path = broken_dir / f"{case_id}.xlsx"
    # cross_sheet_error 把某公式的跨表引用换成另一张表；若换到的表恰好
    # 引用回目标所在链（P&L!H2 引用 Dashboard!H2 时会经 Dashboard!H2 =
    # 'P&L'!H4 绕回），会意外制造依赖循环，案例退化成循环案例。这里
    # 搜索一个"无循环、5 个目标互异且候选集恰为注入目标"的种子。
    injected = None
    for cross_sheet_seed in range(1, 200):
        if broken_path.exists():
            broken_path.unlink()
        specs = [
            (inject_wrong_range, 11),
            (inject_wrong_operator, 7),
            (inject_missing_formula, 81),
            (inject_cross_sheet_error, cross_sheet_seed),
            (inject_wrong_cell_reference, 42),
        ]
        try:
            candidates = inject_faults(str(gold_path), str(broken_path), specs)
        except RuntimeError:
            continue  # 目标冲突，换种子
        index = parse_workbook(str(broken_path))
        graph = DependencyGraph(index)
        graph.build()
        if graph.cycle_sccs():
            continue
        try:
            _assert_candidates(broken_path, {c.target_cell for c in candidates})
        except AssertionError:
            continue
        injected = candidates
        break
    if injected is None:
        raise RuntimeError("hard_004: 200 个种子内找不到无循环的组合")
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            {
                "target_cell": case.target_cell,
                "error_type": case.error_type,
                "old_formula": case.old_formula,
                "gold_formula": case.gold_formula,
            }
            for case in injected
        ],
        "expected_status": "success",
        "static_candidate_count": len(injected),
    })

    # ── hard_005 双断链 ──────────────────────────────────────────────
    case_id = "hard_005"
    broken_path = broken_dir / f"{case_id}.xlsx"
    _copy_gold(gold_path, broken_path)
    wb = load_workbook(broken_path)
    # 两条上游链各自断一处，汇入同一下游 P&L!F2；Agent 必须两处都修，
    # 下游级联验证在上游未齐时为 not_applicable。
    wb["Revenue"]["F2"] = "=RawData!F2+RawData!F3"
    wb["Cost"]["F2"] = "=RawData!F2*RawData!F4-RawData!F5"
    wb.save(broken_path)
    _assert_candidates(broken_path, {"Revenue!F2", "Cost!F2"})
    gold_formulas = _gold_formulas(
        gold_path, [("Revenue", "F2"), ("Cost", "F2")],
    )
    records.append({
        "case_id": case_id,
        "gold_path": str(gold_path),
        "broken_path": str(broken_path),
        "errors": [
            _error(gold_formulas, "Revenue!F2", "wrong_operator",
                   "=RawData!F2+RawData!F3"),
            _error(gold_formulas, "Cost!F2", "wrong_operator",
                   "=RawData!F2*RawData!F4-RawData!F5"),
        ],
        "expected_status": "success",
        "static_candidate_count": 2,
    })

    output = Path(output_json)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    return records


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[2]
    gold = root / "evaluation/data/gold/gold.xlsx"
    cases = build_hard_cases(
        str(gold),
        str(root / "evaluation/data/broken_hard"),
        str(root / "evaluation/data/hard_cases.json"),
    )
    print(f"hard cases: {len(cases)}")
    for case in cases:
        print(
            f"  {case['case_id']}: errors={len(case['errors'])} "
            f"status={case['expected_status']} "
            f"dismissed={case.get('expected_dismissed', '-')} "
            f"skipped={case.get('expected_skipped', '-')} "
            f"deferred={case.get('expected_deferred', '-')}"
        )
