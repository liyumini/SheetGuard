"""用于确定性工作簿调查的 LangChain 工具。"""
from __future__ import annotations

from collections import defaultdict

from networkx.exception import NetworkXError
from langchain_core.tools import BaseTool, tool
from openpyxl.utils import coordinate_to_tuple, get_column_letter

from sheetguard.spreadsheet.dependency_graph import DependencyGraph
from sheetguard.spreadsheet.formula_parser import to_relative_template
from sheetguard.spreadsheet.model import WorkbookIndex


def make_investigation_tools(
    index: WorkbookIndex, graph: DependencyGraph
) -> list[BaseTool]:
    """创建绑定到单个工作簿索引和依赖图的调查工具。

    返回的闭包不依赖全局工作簿状态，因此调用方可以安全地
    为不同工作簿创建各自的工具集合。
    """

    @tool
    def list_sheets() -> str:
        """列出所有工作表名称，以及各表的公式和常量单元格数量。"""
        lines = []
        for sheet in index.sheets:
            formula_count = sum(1 for cell in sheet.cells.values() if cell.is_formula)
            constant_count = len(sheet.cells) - formula_count
            lines.append(
                f"{sheet.name}: {len(sheet.cells)} cells "
                f"({formula_count} formulas, {constant_count} constants)"
            )
        return "\n".join(lines) if lines else "No sheets found."

    @tool
    def get_sheet_summary(sheet: str) -> str:
        """获取一个工作表的单元格、公式和常量数量统计。"""
        sheet_info = index._sheet_by_name(sheet)
        if sheet_info is None:
            return f"Sheet '{sheet}' not found."
        formula_count = sum(
            1 for cell in sheet_info.cells.values() if cell.is_formula
        )
        return (
            f"Sheet '{sheet}': {len(sheet_info.cells)} total cells, "
            f"{formula_count} formula cells, "
            f"{len(sheet_info.cells) - formula_count} constant cells."
        )

    @tool
    def read_formula(cell: str) -> str:
        """读取 ``Sheet!A1`` 的公式或常量值。"""
        info = index.cell(cell)
        if info is None:
            return f"Cell '{cell}' not found."
        if info.is_formula:
            return f"Formula: {info.formula}"
        return f"Constant: {info.value}"

    @tool
    def inspect_range(sheet: str, start_cell: str, end_cell: str) -> str:
        """检查一个矩形单元格范围，最多 20 行 × 20 列。"""
        sheet_info = index._sheet_by_name(sheet)
        if sheet_info is None:
            return f"Sheet '{sheet}' not found."
        try:
            start_row, start_col = coordinate_to_tuple(start_cell)
            end_row, end_col = coordinate_to_tuple(end_cell)
        except ValueError:
            return f"Invalid cell range: {start_cell}:{end_cell}."
        if start_row > end_row or start_col > end_col:
            return f"Invalid cell range: {start_cell}:{end_cell}."

        rows = []
        for row in range(start_row, min(end_row, start_row + 19) + 1):
            values = []
            for col in range(start_col, min(end_col, start_col + 19) + 1):
                address = f"{get_column_letter(col)}{row}"
                info = sheet_info.cells.get(address)
                if info is None:
                    values.append("")
                elif info.is_formula:
                    values.append(info.formula or "")
                else:
                    values.append(str(info.value))
            rows.append(" | ".join(values))
        return "\n".join(rows) if rows else "(empty range)"

    @tool
    def compare_neighbor_formulas(
        cell: str, axis: str = "row", window: int = 3
    ) -> str:
        """将公式的归一化模式与附近的公式进行比较。

        ``axis`` 为 ``row`` 表示水平邻居，``col`` 表示垂直邻居。
        每侧最多检查请求的 window 个邻居。
        """
        info = index.cell(cell)
        if info is None:
            return f"Cell '{cell}' not found."
        if not info.is_formula or not info.formula:
            return f"Cell '{cell}' is not a formula cell."
        if axis not in {"row", "col"}:
            return "Invalid axis; use 'row' or 'col'."
        if window < 0:
            return "Invalid window; use a non-negative integer."

        host_template = to_relative_template(info.formula, info.ref)
        lines = [
            f"Host ({cell}): {info.formula}",
            f"Template: {host_template}",
            f"Neighbors ({axis}):",
        ]
        sheet_info = index._sheet_by_name(info.ref.sheet)
        if sheet_info is None:
            return f"Sheet '{info.ref.sheet}' not found."

        neighbors = []
        for offset in range(-window, window + 1):
            if offset == 0:
                continue
            row = info.ref.row + offset if axis == "col" else info.ref.row
            col = info.ref.col if axis == "col" else info.ref.col + offset
            if row < 1 or col < 1:
                continue
            address = f"{get_column_letter(col)}{row}"
            neighbor = sheet_info.cells.get(address)
            if neighbor and neighbor.is_formula and neighbor.formula:
                template = to_relative_template(neighbor.formula, neighbor.ref)
                marker = "✓" if template == host_template else "✗"
                neighbors.append(f"  {marker} {address}: {neighbor.formula}")

        lines.extend(neighbors or ["  (no formula neighbors)"])
        return "\n".join(lines)

    @tool
    def find_formula_pattern(sheet: str) -> str:
        """查找归一化后的公式模式，以及使用每种模式的单元格。"""
        sheet_info = index._sheet_by_name(sheet)
        if sheet_info is None:
            return f"Sheet '{sheet}' not found."

        patterns: dict[str, list[str]] = defaultdict(list)
        for info in sheet_info.cells.values():
            if info.is_formula and info.formula:
                patterns[to_relative_template(info.formula, info.ref)].append(
                    info.ref.address
                )

        lines = [f"Formula patterns in '{sheet}':"]
        for template, addresses in sorted(
            patterns.items(), key=lambda item: (-len(item[1]), item[0])
        ):
            lines.append(f"  Pattern ({len(addresses)} cells): {template}")
            if len(addresses) <= 5:
                lines.append(f"    Cells: {', '.join(addresses)}")
        return "\n".join(lines)

    @tool
    def trace_precedents(cell: str, depth: int = 1) -> str:
        """追踪 ``cell`` 所依赖的单元格。"""
        try:
            precedents = graph.precedents_of(cell, depth=depth)
        except NetworkXError:
            return f"Cell '{cell}' was not found in the dependency graph."
        except (TypeError, ValueError):
            return f"Invalid dependency query for '{cell}'."
        return f"Precedents of {cell}: {precedents or 'none'}"

    @tool
    def trace_dependents(cell: str, depth: int = 1) -> str:
        """追踪依赖 ``cell`` 的单元格。"""
        try:
            dependents = graph.dependents_of(cell, depth=depth)
        except NetworkXError:
            return f"Cell '{cell}' was not found in the dependency graph."
        except (TypeError, ValueError):
            return f"Invalid dependency query for '{cell}'."
        return f"Dependents of {cell}: {dependents or 'none'}"

    @tool
    def get_dependency_path(src: str, dst: str) -> str:
        """查找从 ``src`` 到 ``dst`` 的依赖路径。"""
        path = graph.path(src, dst)
        if path:
            return " → ".join(path)
        return f"No dependency path from {src} to {dst}."

    return [
        list_sheets,
        get_sheet_summary,
        read_formula,
        inspect_range,
        compare_neighbor_formulas,
        find_formula_pattern,
        trace_precedents,
        trace_dependents,
        get_dependency_path,
    ]
