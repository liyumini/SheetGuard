"""终端聚合格（孤立的行/列汇总公式）的识别判据。

StaticInspector 用它决定哪些单元格豁免"家族少数派"报警，
Verifier 用它决定哪些修复目标豁免"邻居模式一致性"检查。
两处共用同一份判据，避免规则各自演化后互相矛盾。
"""
from __future__ import annotations

import re
from collections import Counter

from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils import get_column_letter, range_boundaries

from sheetguard.spreadsheet.formula_parser import extract_refs, to_relative_template
from sheetguard.spreadsheet.model import CellInfo, CellRef, WorkbookIndex

# 视为“终端聚合”的函数名：这类公式对相邻明细做汇总，模板与明细不同是正常的。
AGGREGATE_FUNCTIONS = (
    "SUM", "AVERAGE", "COUNT", "COUNTA", "MAX", "MIN", "PRODUCT", "SUBTOTAL",
    "SUMIF", "COUNTIF",
)

_AGGREGATE_RE = re.compile(r"(?:%s)\s*\(" % "|".join(AGGREGATE_FUNCTIONS))


def has_aggregate_function(formula: str) -> bool:
    """公式是否包含聚合函数调用（如 SUM/AVERAGE/...）。"""
    return bool(formula and _AGGREGATE_RE.search(formula.upper()))


# 相对模板的引用原子：可选的跨表前缀 + 行列位移。
_TEMPLATE_ATOM = r"(?:S\([^)]*\))?(?:RR?\([^)]*\))"


def normalize_template(template: str) -> str:
    """把相对模板归一化到规范形：去空白、化简一元负括号。

    ``=Revenue!E2 - Cost!E2``（多空格）与 ``=Revenue!H2+(-Cost!H2)``
    （一元取负）在 Excel 语义上与主流模板等价，但不归一化会生成不同
    的模板串——09-12 hard 实验中 2 次正确提案因此被 Verifier 的
    pattern 检查误拒。家族主流模板与候选模板统一走本函数，保证
    比较口径一致（与 formula_equivalent 评测器的语义等价判定对齐）。
    """
    if not template:
        return template or ""
    t = re.sub(r"\s+", "", template)
    prev = None
    while prev != t:
        prev = t
        # 一元负括号 (-X) -> -X；一元正/冗余括号 (+X)/(X) -> X。
        t = re.sub(r"\(-(" + _TEMPLATE_ATOM + r")\)", r"-\1", t)
        t = re.sub(r"\(\+?(" + _TEMPLATE_ATOM + r")\)", r"\1", t)
        # 表达式开头的一元正号去掉（一元负号保留，它承载语义）。
        t = re.sub(r"^\+", "", t)
        # 相邻一元符号化简：++/-- -> +，+-/-+ -> -。
        t = re.sub(
            r"([+\-])[+\-]",
            lambda m: "+" if m.group(0) in ("++", "--") else "-",
            t,
        )
    return t


# 相对模板中的范围原子：RR(行1偏移, 列1偏移, 行2偏移, 列2偏移)。
# SUMRR(...) 这类"函数名紧贴范围原子"的拼接同样以 RR( 开头匹配。
_RANGE_ATOM_RE = re.compile(r"RR\((-?\d+),(-?\d+),(-?\d+),(-?\d+)\)")


def structural_template(template: str) -> str:
    """把相对模板中的范围原子归一化为尺寸，得到结构模板。

    同文复制的家族（如 12 格相同的 ``=VLOOKUP(..., Rate!A2:B13, ...)``）
    因绝对查找区逐宿主相对化而产生互不相同的模板——相对模板层看不出
    "写法相同"。把 RR(行1,列1,行2,列2) 归一为尺寸 RR[H×W] 后，结构
    相同的公式重新聚成主流。单引用原子（R(行,列)）保留位移：家族平移
    语义（明细行逐列右移）由它承载，丢弃会把真实错位放行。
    """
    def _shape(match: "re.Match[str]") -> str:
        r1, c1, r2, c2 = (int(group) for group in match.groups())
        return f"RR[{r2 - r1 + 1}x{c2 - c1 + 1}]"

    return _RANGE_ATOM_RE.sub(_shape, template)


def aggregate_span(index: WorkbookIndex, info: CellInfo) -> tuple[str, tuple[int, int, int, int]] | None:
    """返回孤立聚合公式的（方向，范围），不是一维本地聚合时返回 None。

    聚合范围必须恰好包含一个本地引用（如 =SUM(B2:M2)）、沿聚合方向
    与聚合单元格同行或同列，且不包含聚合单元格自身。二维范围（如
    整表总计）或跨表聚合不属于这里定义的"终端聚合格"。
    """
    if not info.formula or not has_aggregate_function(info.formula):
        return None
    try:
        refs = extract_refs(info.formula)
    except (IndexError, TokenizerError, TypeError, ValueError):
        return None
    if len(refs) != 1:
        return None
    sheet, reference = refs[0]
    if sheet is not None or ":" not in reference:
        return None
    try:
        boundaries = range_boundaries(reference)
    except (IndexError, TokenizerError, TypeError, ValueError):
        return None
    if not all(type(value) is int and value > 0 for value in boundaries):
        return None
    min_col, min_row, max_col, max_row = boundaries
    if min_row == max_row == info.ref.row and not min_col <= info.ref.col <= max_col:
        return ("row", boundaries)
    if min_col == max_col == info.ref.col and not min_row <= info.ref.row <= max_row:
        return ("col", boundaries)
    return None


def is_wellformed_aggregate(index: WorkbookIndex, info: CellInfo) -> bool:
    """判断孤立聚合公式是否完整覆盖了相邻的公式连续段。

    合法的行/列汇总（如 =SUM(B2:M2)）应该"正好盖住"旁边的公式：
    聚合范围两端外侧沿聚合轴紧贴的单元格都不是公式。外侧还有公式
    说明范围没有盖全，这本身就是 wrong_range 的特征。
    """
    span = aggregate_span(index, info)
    if span is None:
        return False
    axis, (min_col, min_row, max_col, max_row) = span
    sheet = index._sheet_by_name(info.ref.sheet)
    if sheet is None:
        return False
    # 除了跨度两端外侧紧贴的格子，还要检查目标靠外一侧“再往外一格”：
    # 例如 I2 =SUM(B2:H2) 时，紧贴格 I2 就是目标自身会被跳过，但右侧
    # J2 仍有公式说明范围没盖全，不能享受“终端汇总”豁免。
    if axis == "row":
        outside_cells = [
            (min_col - 1, min_row),
            (min_col - 2, min_row),
            (max_col + 1, max_row),
            (max_col + 2, max_row),
        ]
    else:
        outside_cells = [
            (min_col, min_row - 1),
            (min_col, min_row - 2),
            (min_col, max_row + 1),
            (min_col, max_row + 2),
        ]
    for col, row in outside_cells:
        if col < 1 or row < 1:
            continue
        if (col, row) == (info.ref.col, info.ref.row):
            continue
        neighbor = sheet.cells.get(f"{get_column_letter(col)}{row}")
        if neighbor is not None and neighbor.is_formula and neighbor.formula:
            return False
    return True


def _axis_formula_families(index: WorkbookIndex, ref: CellRef) -> tuple[list[CellInfo], list[CellInfo]]:
    """返回目标所在工作表上同一行、同一列的公式单元格（目标自身各计入一份）。

    与 StaticInspector._family_groups 的归属方式一致：一个公式单元格同时
    属于其行家族与列家族；目标单元格也计入两侧，宿主家族按成员数取多的一侧。
    """
    sheet = index._sheet_by_name(ref.sheet)
    rows: list[CellInfo] = []
    cols: list[CellInfo] = []
    if sheet is None:
        return rows, cols
    for info in sheet.cells.values():
        if not info.is_formula or not info.formula:
            continue
        if info.ref.full_address == ref.full_address:
            rows.append(info)
            cols.append(info)
        elif info.ref.row == ref.row:
            rows.append(info)
        elif info.ref.col == ref.col:
            cols.append(info)
    return rows, cols


def _cross_sheet_set(info: CellInfo) -> frozenset[str] | None:
    """返回公式引用的跨表名称集合；格式错误时返回 None。"""
    if not info.formula:
        return None
    try:
        return frozenset(sheet for sheet, _ in extract_refs(info.formula) if sheet is not None)
    except (IndexError, TokenizerError, TypeError, ValueError):
        return None


def family_context_data(index: WorkbookIndex, ref: CellRef, max_siblings: int = 3) -> dict:
    """返回目标所在行/列家族的修复上下文（目标自身不列入兄弟）。

    返回字段：
      dominant_template: 家族中出现次数 >= 2 的主流相对模板；没有则 None
      dominant_structural_template: 范围原子按尺寸归一化的结构主流模板，
        同样要求出现 >= 2 次；相对模板层互异但结构同构的同文复制家族
        （无 $ 锚定的固定查找区）由它兜底；没有则 None
      siblings         : 与主流模板一致的兄弟公式原文（"地址: 公式"），最多 max_siblings 条
      dominant_sheets  : 家族公式中最常见的跨表引用集合（排序后的列表）

    StaticInspector.family_context（喂给修复模型的证据）与 Verifier 的
    pattern 检查（验证候选公式）共用同一份判据，避免规则各自演化。
    """
    rows, cols = _axis_formula_families(index, ref)
    target_info = index.cell(ref.full_address)
    target_is_aggregate = (
        target_info is not None
        and target_info.is_formula
        and target_info.formula
        and has_aggregate_function(target_info.formula)
    )
    if target_is_aggregate:
        # 目标是汇总格（如行尾合计）时，优先选以聚合公式为主的轴（合计列），
        # 而不是成员更多的明细行——明细行模板与汇总格写法不同，会带偏修复。
        row_agg = sum(1 for m in rows if has_aggregate_function(m.formula))
        col_agg = sum(1 for m in cols if has_aggregate_function(m.formula))
        if row_agg or col_agg:
            host = rows if row_agg >= col_agg else cols
        else:
            host = rows if len(rows) >= len(cols) else cols
    else:
        host = rows if len(rows) >= len(cols) else cols
    if len(host) < 3:
        return {
            "dominant_template": None,
            "dominant_structural_template": None,
            "siblings": [],
            "dominant_sheets": [],
        }
    templates: list[tuple[CellInfo, str]] = []
    for member in host:
        try:
            templates.append(
                (member, normalize_template(to_relative_template(member.formula, member.ref)))
            )
        except (IndexError, TokenizerError, TypeError, ValueError):
            continue
    dominant_template: str | None = None
    dominant_structural_template: str | None = None
    siblings: list[str] = []
    dominant_sheets: list[str] = []
    if not templates:
        return {
            "dominant_template": None,
            "dominant_structural_template": None,
            "siblings": [],
            "dominant_sheets": [],
        }
    counts = Counter(template for _, template in templates)
    dominant_template, dominant_count = counts.most_common(1)[0]
    if dominant_count < 2:
        # count=1 的模板没有多数派支撑（完整异构家族里最常见串也只出现
        # 一次）；把它当主流会把 gold 公式结构性判死（complete_018 的
        # Revenue!L4）。与 StaticInspector._signal_pattern_anomaly 的
        # dominant_count >= 2 门槛保持一致。
        dominant_template = None
    structural_counts = Counter(structural_template(t) for _, t in templates)
    dominant_structural_template, structural_count = structural_counts.most_common(1)[0]
    if structural_count < 2:
        dominant_structural_template = None
    if dominant_count >= 2:
        for member, template in templates:
            if len(siblings) >= max_siblings:
                break
            if member.ref.full_address == ref.full_address:
                continue
            if template == dominant_template:
                siblings.append(f"{member.ref.full_address}: {member.formula}")
    sheet_counts = Counter(
        sheet_set
        for member, _ in templates
        if (sheet_set := _cross_sheet_set(member)) is not None
    )
    if sheet_counts:
        dominant_sheets = sorted(sheet_counts.most_common(1)[0][0])
    return {
        "dominant_template": dominant_template,
        "dominant_structural_template": dominant_structural_template,
        "siblings": siblings,
        "dominant_sheets": dominant_sheets,
    }


# 公开别名：formula_expectation 的视角预测器与 StaticInspector 家族上下文
# 共用同一份行/列家族归属判据，避免两处规则各自演化（设计 4.7）。
axis_formula_families = _axis_formula_families
