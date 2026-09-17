"""多视角公式期望推断（v1.6 设计 4）。

StaticInspector 的 dependency_anomaly 从"先选唯一 family，再看目标格是否
少数派"改为"先把目标格遮住，让横向/纵向/依赖等多个结构视角独立预测
这个位置本该是什么，再用高置信共识与实际公式比较"（设计 4.1/4.3）。

核心概念：
- leave-one-out：目标格自身不参与任何支持集的建立，避免错误公式污染
  基准（self-contamination，设计 4.1）。
- 多视角独立预测：每个视角输出一个 Prediction；共识只来自"相互一致且
  强度足够"的视角，冲突或支持不足时拒绝判断（no reliable expectation）。
- 公式依赖 ≠ 修复前置条件：本模块只产出结构证据；是否升级为修复前置
  由调度器在 Phase 2 按 candidate confidence 决定（设计 5.1）。
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass

from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils import get_column_letter, range_boundaries

from sheetguard.spreadsheet.aggregates import (
    axis_formula_families,
    has_aggregate_function,
    is_wellformed_aggregate,
)
from sheetguard.spreadsheet.formula_parser import extract_refs, to_relative_template
from sheetguard.spreadsheet.model import CellInfo, CellRef, WorkbookIndex

# 视角预测进入共识的最低置信度（设计 4.5 的 view_min_confidence）。
VIEW_MIN_CONFIDENCE = 0.5
# 共识成立的最低置信度（设计 4.5 的 minimum_consensus_confidence）。
MINIMUM_CONSENSUS_CONFIDENCE = 0.5
# 共识只剩单一视角时按此系数折减：聚合视角 0.9*0.6=0.54 仍可成立，
# 而单支持者的依赖视角 0.75*0.6=0.45 被拒绝——孤立弱信号不能驱动强异常。
_SINGLE_VIEW_FACTOR = 0.6
# 同一组内模板互相冲突时按此系数折减（设计 4.3 的冲突降置信）。
_CONFLICT_FACTOR = 0.6


@dataclass(frozen=True)
class FormulaPrediction:
    """一个结构视角对目标格的独立预测（设计 4.2）。"""

    view: str  # horizontal / vertical / dependency
    expected_template: str | None  # 期望相对模板；无模板证据时为 None
    expected_formula: str | None  # 期望公式原文；无模板证据时为 None
    expected_dependencies: frozenset[str]  # 期望跨表引用集合；纯本地引用记 ∅
    confidence: float
    evidence_cells: tuple[str, ...]


@dataclass(frozen=True)
class DeviationEvidence:
    """实际公式与共识期望的偏离证据（设计 4.5 的 deviation）。"""

    is_dependency_outlier: bool
    dependencies_match: bool
    template_matches: bool | None
    expected_dependencies: frozenset[str]
    actual_dependencies: frozenset[str]
    expected_formula: str | None
    expected_template: str | None
    severity: float  # 0.0 一致 / 0.5 仅依赖偏离 / 1.0 依赖与模板同时偏离


@dataclass(frozen=True)
class FormulaExpectation:
    """多视角预测及其共识；confidence < 阈值时表示拒绝判断。"""

    predictions: tuple[FormulaPrediction, ...]
    consensus_template: str | None
    consensus_formula: str | None
    consensus_dependencies: frozenset[str] | None
    confidence: float

    @property
    def has_consensus(self) -> bool:
        return (
            self.confidence >= MINIMUM_CONSENSUS_CONFIDENCE
            and self.consensus_dependencies is not None
        )


def shift_formula(formula: str, row_delta: int, col_delta: int) -> str:
    """把公式里所有引用整体位移 (row_delta, col_delta)，返回新公式文本。

    复制规律假设：同一模板的公式平移后引用同步平移。位移后行列号 < 1
    说明该公式无法合法复制到目标位置，抛 ValueError 由调用方把该支持者
    降级处理。$ 绝对引用与 to_relative_template 同样按普通引用处理。
    """
    if not formula.startswith("="):
        return formula
    parts: list[str] = []
    for tok in Tokenizer(formula).items:
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            value = tok.value
            prefix = ""
            if "!" in value:
                prefix = value[: value.rindex("!") + 1]
                clean = value[value.rindex("!") + 1:]
            else:
                clean = value
            clean = clean.replace("$", "")
            try:
                min_col, min_row, max_col, max_row = range_boundaries(clean)
            except Exception:
                parts.append(value)
                continue
            min_row += row_delta
            max_row += row_delta
            min_col += col_delta
            max_col += col_delta
            if min_row < 1 or min_col < 1:
                raise ValueError("shifted reference out of bounds")
            if (min_col, min_row) == (max_col, max_row):
                parts.append(f"{prefix}{get_column_letter(min_col)}{min_row}")
            else:
                parts.append(
                    f"{prefix}{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{max_row}"
                )
        else:
            parts.append(tok.value)
    # Tokenizer 会剥离前导 =，重建公式文本时补回。
    return "=" + "".join(parts)


def _axis_supporters(
    index: WorkbookIndex,
    ref: CellRef,
    members: list[CellInfo],
    exclude_target: bool,
) -> list[tuple[CellInfo, str]]:
    """把行/列家族成员整理为 (成员, 自身相对模板) 支持者列表。

    exclude_target=True 时跳过目标自身（leave-one-out，检测路径恒为 True；
    False 仅用于诊断对比）。
    """
    supporters: list[tuple[CellInfo, str]] = []
    for member in members:
        if exclude_target and member.ref.full_address == ref.full_address:
            continue
        try:
            template = to_relative_template(member.formula, member.ref)
        except (IndexError, TokenizerError, TypeError, ValueError):
            continue
        supporters.append((member, template))
    return supporters


def _axis_prediction(
    view: str,
    index: WorkbookIndex,
    ref: CellRef,
    members: list[CellInfo],
    exclude_target: bool,
) -> FormulaPrediction | None:
    """横向/纵向视角预测：主流模板组把公式位移到目标位置后期望一致。

    平移不改变相对偏移，同模板组成员对目标位置的期望模板必然相同，
    因此判据就是"主流组内 ≥2 个成员"（单一成员无法区分复制规律与巧合）。
    """
    groups: dict[str, list[CellInfo]] = defaultdict(list)
    for member, template in _axis_supporters(index, ref, members, exclude_target):
        groups[template].append(member)
    ranked = sorted(groups.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    if not ranked or len(ranked[0][1]) < 2:
        return None
    if len(ranked) > 1 and len(ranked[1][1]) == len(ranked[0][1]):
        return None  # 模板组并列（视角冲突）→ 拒绝判断，不强行选 family（设计 4.3）。
    template, group = ranked[0]
    source = min(
        group,
        key=lambda m: (
            abs(m.ref.row - ref.row),
            abs(m.ref.col - ref.col),
            m.ref.full_address,
        ),
    )
    try:
        expected_formula = shift_formula(
            source.formula,
            ref.row - source.ref.row,
            ref.col - source.ref.col,
        )
        expected_deps = frozenset(
            sheet for sheet, _ in extract_refs(expected_formula) if sheet is not None
        )
    except (IndexError, TokenizerError, TypeError, ValueError):
        return None
    size = len(group)
    return FormulaPrediction(
        view=view,
        expected_template=template,
        expected_formula=expected_formula,
        expected_dependencies=expected_deps,
        confidence=min(0.95, 0.75 + 0.05 * size),
        evidence_cells=tuple(sorted(m.ref.full_address for m in group)),
    )


def _dependency_prediction(
    index: WorkbookIndex,
    ref: CellRef,
    exclude_target: bool,
) -> FormulaPrediction | None:
    """依赖视角预测：行/列家族的主流跨表引用集合（设计 4.2）。

    需要严格多于次名分组的支持数（设计 4.3：视角冲突时拒绝判断）；
    top_count=1 时输出 0.75 低置信预测，由共识层决定是否采信。
    """
    rows, cols = axis_formula_families(index, ref)
    counts: dict[frozenset[str], int] = defaultdict(int)
    for member in (*rows, *cols):
        if exclude_target and member.ref.full_address == ref.full_address:
            continue
        try:
            deps = frozenset(
                sheet for sheet, _ in extract_refs(member.formula) if sheet is not None
            )
        except (IndexError, TokenizerError, TypeError, ValueError):
            continue
        counts[deps] += 1
    if not counts:
        return None
    ranked = sorted(counts.items(), key=lambda kv: (-kv[1], sorted(kv[0])))
    top_deps, top_count = ranked[0]
    if len(ranked) > 1 and ranked[1][1] == top_count:
        return None  # 两个同规模依赖集合互相冲突 → 拒绝
    return FormulaPrediction(
        view="dependency",
        expected_template=None,
        expected_formula=None,
        expected_dependencies=top_deps,
        confidence=min(0.9, 0.7 + 0.05 * top_count),
        evidence_cells=(),
    )


def _build_consensus(
    predictions: list[FormulaPrediction],
) -> FormulaExpectation:
    """把多视角预测聚合为共识；冲突或证据不足时返回拒绝判断的期望。

    规则（设计 4.3）：
    - 按期望依赖集合分组，取规模最大的组为共识（并列即冲突 → 拒绝）；
    - 组内只有 1 个视角时按 _SINGLE_VIEW_FACTOR 折减——孤证不驱动强异常；
    - 组内模板互相冲突时丢弃模板共识并按 _CONFLICT_FACTOR 折减，
      但依赖层面的共识仍然可用。
    """
    def _refuse(preds: tuple) -> FormulaExpectation:
        return FormulaExpectation(
            predictions=preds,
            consensus_template=None,
            consensus_formula=None,
            consensus_dependencies=None,
            confidence=0.0,
        )

    frozen = tuple(predictions)
    if not frozen:
        return _refuse(())
    groups: dict[frozenset[str], list[FormulaPrediction]] = defaultdict(list)
    for pred in frozen:
        groups[pred.expected_dependencies].append(pred)
    ranked = sorted(groups.items(), key=lambda kv: (-len(kv[1]), sorted(kv[0])))
    top_deps, top_preds = ranked[0]
    if len(ranked) > 1 and len(ranked[1][1]) == len(top_preds):
        return _refuse(frozen)  # 视角冲突：不强行选唯一 family
    top_size = len(top_preds)
    if top_size == 1:
        confidence = top_preds[0].confidence * _SINGLE_VIEW_FACTOR
    else:
        confidence = min(0.95, 0.75 + 0.05 * top_size)
    templates = {p.expected_template for p in top_preds if p.expected_template}
    if len(templates) > 1:
        confidence *= _CONFLICT_FACTOR
        consensus_template = None
        consensus_formula = None
    elif templates:
        best = max(
            (p for p in top_preds if p.expected_template),
            key=lambda p: p.confidence,
        )
        consensus_template = best.expected_template
        consensus_formula = best.expected_formula
    else:
        consensus_template = None
        consensus_formula = None
    return FormulaExpectation(
        predictions=frozen,
        consensus_template=consensus_template,
        consensus_formula=consensus_formula,
        consensus_dependencies=top_deps,
        confidence=confidence,
    )


def predict_formula_expectation(
    index: WorkbookIndex,
    target: str | CellRef,
    *,
    exclude_target: bool = True,
    target_info: CellInfo | None = None,
) -> FormulaExpectation:
    """对目标格做多视角期望推断（设计 4.5 的新流程）。

    exclude_target=True（leave-one-out，默认）：目标自身不进入任何支持集。
    """
    ref = target if isinstance(target, CellRef) else CellRef.parse(target)
    if target_info is None:
        target_info = index.cell(ref.full_address)
    predictions: list[FormulaPrediction] = []
    rows, cols = axis_formula_families(index, ref)
    for view, members in (("horizontal", rows), ("vertical", cols)):
        pred = _axis_prediction(view, index, ref, members, exclude_target)
        if pred is not None and pred.confidence >= VIEW_MIN_CONFIDENCE:
            predictions.append(pred)
    dep_pred = _dependency_prediction(index, ref, exclude_target)
    if dep_pred is not None and dep_pred.confidence >= VIEW_MIN_CONFIDENCE:
        predictions.append(dep_pred)
    return _build_consensus(predictions)


def _cross_sheet_set(info: CellInfo) -> frozenset[str] | None:
    """返回公式引用的跨表名称集合；公式不可解析时返回 None。"""
    if not info.formula:
        return None
    try:
        return frozenset(
            sheet for sheet, _ in extract_refs(info.formula) if sheet is not None
        )
    except (IndexError, TokenizerError, TypeError, ValueError):
        return None


def compare_actual_to_expectation(
    target: CellInfo,
    expectation: FormulaExpectation,
) -> DeviationEvidence | None:
    """把目标格实际公式与共识期望比较；无共识或公式不可解析时返回 None。"""
    if not expectation.has_consensus:
        return None
    actual = _cross_sheet_set(target)
    if actual is None:
        return None
    dependencies_match = actual == expectation.consensus_dependencies
    template_matches: bool | None = None
    if expectation.consensus_template is not None:
        try:
            template_matches = (
                to_relative_template(target.formula, target.ref)
                == expectation.consensus_template
            )
        except (IndexError, TokenizerError, TypeError, ValueError):
            template_matches = None
    if not dependencies_match and template_matches is False:
        severity = 1.0
    elif not dependencies_match:
        severity = 0.5
    else:
        severity = 0.0
    return DeviationEvidence(
        is_dependency_outlier=not dependencies_match,
        dependencies_match=dependencies_match,
        template_matches=template_matches,
        expected_dependencies=expectation.consensus_dependencies,
        actual_dependencies=actual,
        expected_formula=expectation.consensus_formula,
        expected_template=expectation.consensus_template,
        severity=severity,
    )


def dependency_judgement(
    index: WorkbookIndex,
    info: CellInfo,
) -> tuple[DeviationEvidence | None, FormulaExpectation | None]:
    """单公式格的依赖判定 + 多视角期望（设计 5.2 的 confidence 证据源）。

    聚合格走确定性分支：偏离证据照常返回，期望为 None（无共识概念）。
    其余公式：期望与偏离证据一并返回；has_consensus=False 表示拒绝判断。
    """
    if has_aggregate_function(info.formula or ""):
        return dependency_outlier_evidence(index, info), None
    expectation = predict_formula_expectation(index, info.ref, target_info=info)
    return compare_actual_to_expectation(info, expectation), expectation


def dependency_outlier_evidence(
    index: WorkbookIndex,
    info: CellInfo,
) -> DeviationEvidence | None:
    """单公式格的依赖偏离判定（StaticInspector 的消费入口）。

    聚合格目标（含 SUM/AVERAGE 等聚合函数）走确定性分支：终端汇总格
    应当只引用本表数据（期望依赖 ∅）；带跨表引用且没有盖全相邻公式段
    （非 wellformed 一维本地聚合）即偏离——与 v1.5 的聚合格豁免语义一致，
    避免 Total 被普通复制模板误判（设计 4.2/8.1）。
    其余公式走多视角共识比较：实际依赖与所有高置信共识都不一致才判偏离。
    """
    if has_aggregate_function(info.formula or ""):
        actual = _cross_sheet_set(info)
        if actual is None or not actual:
            return None  # 本地汇总引用 ∅ 是合法的终端聚合格
        if is_wellformed_aggregate(index, info):
            return None  # 盖全的合法汇总豁免（v1.5 语义保持）
        return DeviationEvidence(
            is_dependency_outlier=True,
            dependencies_match=False,
            template_matches=None,
            expected_dependencies=frozenset(),
            actual_dependencies=actual,
            expected_formula=None,
            expected_template=None,
            severity=1.0,
        )
    expectation = predict_formula_expectation(index, info.ref, target_info=info)
    return compare_actual_to_expectation(info, expectation)
