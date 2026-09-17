"""受限公式子集的确定性重算引擎。

支持：+, -, *, /, 一元负号, SUM(范围), AVERAGE/MIN/MAX/COUNT,
IF/SUMIF/COUNTIF/VLOOKUP, 比较运算符, 字符串/布尔常量,
单元格引用, 数字常量。
不支持：数组公式, 复杂嵌套。

空白格语义与 Excel 对齐：聚合函数（AVERAGE/COUNT 及 SUM 的文本规则）
跳过空白格与文本；普通算术把空白格按 0 处理。
"""
from __future__ import annotations
from enum import Enum, auto
from dataclasses import dataclass
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries
from sheetguard.spreadsheet.model import WorkbookIndex, CellRef
from sheetguard.spreadsheet.dependency_graph import DependencyGraph
from sheetguard.spreadsheet.formula_parser import (
    _sanitize_range_str,
    _split_sheet_ref,
    expand_range,
)


class TokType(Enum):
    """公式解析器使用的 Token 类型。

    例如，公式 `=A1+SUM(B1:B3)` 会被拆分为：
    `CELLREF`、`PLUS`、`FUNC`、`LPAREN`、`CELLREF` 等 Token 类型。
    """

    NUMBER = auto()   # 数字，例如 100
    CELLREF = auto()  # 单元格或范围引用，例如 Sheet1!A1、Sheet1!B1:B3
    LPAREN = auto()   # 左括号：(
    RPAREN = auto()   # 右括号：)
    PLUS = auto()     # 加法运算符：+
    MINUS = auto()    # 减法运算符或一元负号：-
    STAR = auto()     # 乘法运算符：*
    SLASH = auto()    # 除法运算符：/
    FUNC = auto()     # 函数名称，例如 SUM
    COMMA = auto()    # 函数参数分隔符：,
    GT = auto()       # 比较：>
    LT = auto()       # 比较：<
    GTE = auto()      # 比较：>=
    LTE = auto()      # 比较：<=
    EQ = auto()       # 比较：=
    NEQ = auto()      # 比较：<>
    STRING = auto()   # 字符串常量，例如 "OK"（value 保留引号原文）
    LOGICAL = auto()  # 布尔字面量：TRUE / FALSE
    EOF = auto()      # Token 流结束标记


@dataclass
class Tok:
    """表示公式中的一个 Token。

    `type` 表示 Token 的类别，`value` 保存原始文本。

    示例：
        `Tok(TokType.NUMBER, "100")` 表示数字 100；
        `Tok(TokType.CELLREF, "Sheet1!A1")` 表示对 Sheet1!A1 的引用；
        `Tok(TokType.PLUS, "+")` 表示加法运算符。
    """

    type: TokType  # Token 类型，例如 NUMBER、CELLREF、PLUS
    value: str     # Token 的具体文本，例如 "100"、"Sheet1!A1"、"+"


class _ExprParser:
    """简单算术表达式和 SUM 表达式的递归下降解析器。

    示例：公式 `A1+B1*2` 会被解析成下面的表达式树，
    其中乘法优先于加法：

            (+)
           /   \\
         A1     (*)
               /   \\
             B1     2

    也就是：`A1 + (B1 * 2)`，而不是 `(A1 + B1) * 2`。
    """

    def __init__(self, tokens: list[Tok]):
        self.tokens = tokens
        self.pos = 0

    def _peek(self) -> Tok:
        """查看当前位置的 Token，但不移动读取位置。

        示例：当前位置是 `+` 时，调用 `_peek()` 会返回 `Tok(PLUS, "+")`，
        再次调用仍然会看到同一个 Token。
        """
        return self.tokens[self.pos] if self.pos < len(self.tokens) else Tok(TokType.EOF, "")

    def _advance(self) -> Tok:
        """读取当前 Token，并将位置向后移动一格。

        示例：当前 Token 是 `NUMBER(10)` 时，调用后返回它，下一次读取会进入
        下一个 Token。
        """
        t = self._peek()
        self.pos += 1
        return t

    def parse(self):
        """从当前 Token 开始解析完整表达式（比较优先级最低）。

        示例：`A1+2>3` 会被解析为 `cmp(GT, binop(+, A1, 2), 3)`。
        """
        return self._parse_cmp()

    def _parse_cmp(self):
        """解析比较层级（> < >= <= = <>），优先级低于四则运算。

        示例：`A1+2 > 3` 解析为 `cmp(>, binop(+, A1, 2), 3)`；
        `A1 > B1 > C1` 左结合为 `cmp(>, cmp(>, A1, B1), C1)`。
        """
        left = self._parse_expr()
        while self._peek().type in (
            TokType.GT, TokType.LT, TokType.GTE,
            TokType.LTE, TokType.EQ, TokType.NEQ,
        ):
            op = self._advance()
            right = self._parse_expr()
            left = ("cmp", op.type, left, right)
        return left

    def _parse_expr(self):
        """解析加法和减法层级。

        示例：`A1+B1-C1` 会解析为 `(A1+B1)-C1`。
        """
        left = self._parse_term()
        while self._peek().type in (TokType.PLUS, TokType.MINUS):
            op = self._advance()
            right = self._parse_term()
            left = ("binop", op.type, left, right)
        return left

    def _parse_term(self):
        """解析乘法和除法层级，优先级高于加减法。

        示例：`A1+B1*C1` 会解析为 `A1+(B1*C1)`，而不是 `(A1+B1)*C1`。
        """
        left = self._parse_factor()
        while self._peek().type in (TokType.STAR, TokType.SLASH):
            op = self._advance()
            right = self._parse_factor()
            left = ("binop", op.type, left, right)
        return left

    def _parse_factor(self):
        """解析最基本的表达式元素。

        支持数字、字符串/布尔常量、单元格引用、括号、一元负号和函数。
        示例：`10`、`"OK"`、`TRUE`、`A1`、`(A1+1)`、`-A1`、`SUM(A1:A3)`。
        """
        tok = self._peek()
        if tok.type == TokType.NUMBER:
            self._advance()
            return ("number", tok.value)
        elif tok.type == TokType.STRING:
            self._advance()
            return ("string", tok.value[1:-1])
        elif tok.type == TokType.LOGICAL:
            self._advance()
            return ("bool", tok.value.upper() == "TRUE")
        elif tok.type == TokType.CELLREF:
            self._advance()
            return ("cellref", tok.value)
        elif tok.type == TokType.LPAREN:
            self._advance()
            expr = self._parse_cmp()
            if self._peek().type != TokType.RPAREN:
                raise ValueError(f"Expected ) at position {self.pos}")
            self._advance()
            return expr
        elif tok.type == TokType.MINUS:
            self._advance()
            return ("unary", "-", self._parse_factor())
        elif tok.type == TokType.FUNC:
            name = self._advance().value
            if self._peek().type != TokType.LPAREN:
                raise ValueError(f"Expected ( after function {name}")
            self._advance()  # consume (
            args = []
            if self._peek().type != TokType.RPAREN:
                args.append(self._parse_cmp())
                while self._peek().type == TokType.COMMA:
                    self._advance()
                    args.append(self._parse_cmp())
            if self._peek().type != TokType.RPAREN:
                raise ValueError(f"Expected ) after function args at {self.pos}")
            self._advance()  # consume )
            return ("func", name, args)
        else:
            raise ValueError(f"Unexpected token {tok.type} ({tok.value}) at pos {self.pos}")


def _tokenize_openpyxl(formula: str, host: CellRef) -> list[Tok]:
    """将 openpyxl Token 转换为项目内部的简化 Token 流。

    同时把单元格引用补全为完整地址，并映射比较运算符、
    字符串/布尔常量（支持：…IF/SUMIF/COUNTIF/VLOOKUP, 比较运算符,
    字符串/布尔常量）。
    示例：当前单元格在 `P&L!C6` 时，公式中的 `B2` 会转换为
    `Tok(CELLREF, "P&L!B2")`；`">100"` 会转换为 `Tok(STRING, '">100"')`。
    """
    expr = formula  # Tokenizer handles the "=" prefix
    tokens: list[Tok] = []
    for tok in Tokenizer(expr).items:
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            sheet, rest = _split_sheet_ref(tok.value)
            rest = _sanitize_range_str(rest)
            sheet_name = sheet or host.sheet
            tokens.append(Tok(TokType.CELLREF, f"{sheet_name}!{rest}"))
        elif tok.type == "OPERAND" and tok.subtype == "NUMBER":
            tokens.append(Tok(TokType.NUMBER, tok.value))
        elif tok.type == "OPERAND" and tok.subtype == "TEXT":
            tokens.append(Tok(TokType.STRING, tok.value))
        elif tok.type == "OPERAND" and tok.subtype == "LOGICAL":
            tokens.append(Tok(TokType.LOGICAL, tok.value))
        elif tok.type == "FUNC":
            if tok.subtype == "OPEN":
                name = tok.value.rstrip("(")
                tokens.append(Tok(TokType.FUNC, name))
                tokens.append(Tok(TokType.LPAREN, "("))
            elif tok.subtype == "CLOSE":
                tokens.append(Tok(TokType.RPAREN, ")"))
        elif tok.type == "PAREN":
            if tok.subtype == "OPEN":
                tokens.append(Tok(TokType.LPAREN, "("))
            elif tok.subtype == "CLOSE":
                tokens.append(Tok(TokType.RPAREN, ")"))
        elif tok.type == "OPERATOR-INFIX":
            if tok.value == "+":
                tokens.append(Tok(TokType.PLUS, "+"))
            elif tok.value == "-":
                tokens.append(Tok(TokType.MINUS, "-"))
            elif tok.value == "*":
                tokens.append(Tok(TokType.STAR, "*"))
            elif tok.value == "/":
                tokens.append(Tok(TokType.SLASH, "/"))
            elif tok.value in (">", "<", ">=", "<=", "=", "<>"):
                op_type = {
                    ">": TokType.GT, "<": TokType.LT,
                    ">=": TokType.GTE, "<=": TokType.LTE,
                    "=": TokType.EQ, "<>": TokType.NEQ,
                }[tok.value]
                tokens.append(Tok(op_type, tok.value))
        elif tok.type == "OPERATOR-PREFIX":
            if tok.value == "-":
                tokens.append(Tok(TokType.MINUS, "-"))
        elif tok.type == "SEP" and tok.value == ",":
            tokens.append(Tok(TokType.COMMA, ","))
    return tokens


def validate_formula(formula: str, host: CellRef) -> None:
    """验证 V1 支持的公式语法，但不执行计算。

    示例：`=A1+B1` 可以通过；`=A1+` 会抛出语法异常。
    函数名称即使语法正确，若不是支持的函数，也会在重算阶段被拒绝。
    """
    tokens = _tokenize_openpyxl(formula, host)
    parser = _ExprParser(tokens)
    parser.parse()
    if parser._peek().type is not TokType.EOF:
        raise ValueError(f"Unexpected token after formula at position {parser.pos}")


# 重算引擎支持的函数全集（_eval_ast 的 func 分支）。proposal 预检用它
# 把"不支持的函数"挡在 attempt 副本之前——否则这些公式到重算才变成
# NaN，反馈里只剩"无法重算"，模型不知道错在哪类（complete_010 的 ROUND）。
SUPPORTED_FUNCTIONS = frozenset(
    {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "IF", "SUMIF", "COUNTIF", "VLOOKUP"}
)


def unsupported_functions(formula: str) -> list[str]:
    """返回公式中重算引擎不支持的函数名列表（大小写不敏感、保序去重）。"""
    names: list[str] = []
    for tok in Tokenizer(formula).items:
        if tok.type == "FUNC" and tok.subtype == "OPEN":
            name = tok.value.rstrip("(").upper()
            if name not in SUPPORTED_FUNCTIONS and name not in names:
                names.append(name)
    return names


def _expand_range_values(full_addr: str, values: dict[str, float | object]) -> list[float]:
    """展开 SUM 范围并读取其中的数字值，忽略文本常量。

    示例：`A1:A3` 的值为 `10、"文本"、20` 时，返回 `[10, 20]`，
    与 Excel 的 `SUM(A1:A3)` 行为一致。空白格（_BLANK）对求和而言
    等价于 0，无须显式排除。
    """
    cells = _range_cells(full_addr)
    return [
        value for s, a in cells
        if isinstance((value := values.get(f"{s}!{a}", 0.0)), (int, float))
    ]


def _eval_ast(ast, values: dict[str, float | object], index: WorkbookIndex | None = None) -> float:
    """递归计算抽象语法树 AST。

    示例：公式 `A1+B1*2` 且 A1=10、B1=5 时，按运算优先级返回 20.0。
    支持单元格、四则运算、一元负号、SUM 与 AVERAGE/MIN/MAX/COUNT。

    ``index`` 供聚合函数区分"空白格"（未写入或值为 None 的格子不计入
    AVERAGE/COUNT，与 Excel 一致）；不可用时退化为 values 字典语义。
    """
    kind = ast[0]
    if kind == "number":
        return float(ast[1])
    elif kind == "cellref":
        value = values.get(ast[1], 0.0)
        # 空白格参与算术时按 0 处理（Excel 语义），仅聚合函数跳过它。
        return 0.0 if value is _BLANK else value
    elif kind == "binop":
        _, op, left, right = ast
        lv = _eval_ast(left, values, index)
        rv = _eval_ast(right, values, index)
        if op == TokType.PLUS:
            return lv + rv
        elif op == TokType.MINUS:
            return lv - rv
        elif op == TokType.STAR:
            return lv * rv
        elif op == TokType.SLASH:
            if rv == 0:
                raise ZeroDivisionError("formula division by zero")
            return lv / rv
        else:
            raise ValueError(f"Unknown operator: {op}")
    elif kind == "unary":
        return -_eval_ast(ast[2], values, index)
    elif kind == "cmp":
        _, op, left, right = ast
        lv = _numeric_of(_eval_ast(left, values, index))
        rv = _numeric_of(_eval_ast(right, values, index))
        if op == TokType.GT:
            return 1.0 if lv > rv else 0.0
        if op == TokType.LT:
            return 1.0 if lv < rv else 0.0
        if op == TokType.GTE:
            return 1.0 if lv >= rv else 0.0
        if op == TokType.LTE:
            return 1.0 if lv <= rv else 0.0
        if op == TokType.EQ:
            return 1.0 if lv == rv else 0.0
        if op == TokType.NEQ:
            return 1.0 if lv != rv else 0.0
        raise ValueError(f"Unknown comparison operator: {op}")
    elif kind == "string":
        return _TEXT_VALUE
    elif kind == "bool":
        return 1.0 if ast[1] else 0.0
    elif kind == "func":
        name = ast[1].upper()
        if name == "SUM":
            total = 0.0
            for arg in ast[2]:
                if arg[0] == "cellref" and ":" in arg[1]:
                    # Expand range reference; text members are ignored.
                    total += sum(_expand_range_values(arg[1], values))
                elif arg[0] == "cellref":
                    # Excel SUM ignores a direct text reference as well.
                    value = values.get(arg[1], 0.0)
                    if value is not _TEXT_VALUE and value is not _BLANK:
                        total += value
                else:
                    total += _eval_ast(arg, values, index)
            return total
        elif name in ("AVERAGE", "MIN", "MAX", "COUNT"):
            numbers: list[float] = []
            for arg in ast[2]:
                if arg[0] == "cellref":
                    numbers.extend(_cell_or_range_numerics(arg[1], values, index))
                else:
                    numbers.append(float(_eval_ast(arg, values, index)))
            if name == "COUNT":
                return float(len(numbers))
            if not numbers:
                if name == "AVERAGE":
                    # Excel：无可参与计算的数值 → #DIV/0! → 这里记 NaN。
                    raise ZeroDivisionError("AVERAGE over no numeric values")
                return 0.0  # Excel：空集的 MIN/MAX 返回 0
            if name == "AVERAGE":
                return sum(numbers) / len(numbers)
            if name == "MIN":
                return min(numbers)
            return max(numbers)
        elif name == "IF":
            if len(ast[2]) not in (2, 3):
                raise ValueError(
                    f"IF requires 2 or 3 arguments, got {len(ast[2])}"
                )
            cond = _numeric_of(_eval_ast(ast[2][0], values, index))
            if cond != 0.0:
                return _eval_ast(ast[2][1], values, index)
            if len(ast[2]) == 3:
                return _eval_ast(ast[2][2], values, index)
            return 0.0  # Excel：条件为 FALSE 且无第三参 → FALSE → 0
        elif name in ("SUMIF", "COUNTIF"):
            if len(ast[2]) not in (2, 3) or name == "COUNTIF" and len(ast[2]) != 2:
                raise ValueError(f"{name} argument count invalid: {len(ast[2])}")
            range_arg = ast[2][0]
            if range_arg[0] != "cellref" or ":" not in range_arg[1]:
                raise ValueError(f"{name} range must be a range reference")
            criteria = _parse_criteria(ast[2][1], values, index)
            if name == "SUMIF" and len(ast[2]) == 3:
                sum_arg = ast[2][2]
                if sum_arg[0] != "cellref" or ":" not in sum_arg[1]:
                    raise ValueError("SUMIF sum_range must be a range reference")
                sum_cells = expand_range(*_split_sheet_ref(sum_arg[1]))
            else:
                sum_cells = None
            cells = _range_cells(range_arg[1])
            if sum_cells is not None and len(sum_cells) != len(cells):
                raise ValueError("SUMIF sum_range must match range size")
            sum_cells = sum_cells if sum_cells is not None else cells
            matched = 0
            total = 0.0
            for (rs, ra), (vs, va) in zip(cells, sum_cells):
                cell_value = values.get(f"{rs}!{ra}", _BLANK)
                hit = _criteria_match(cell_value, criteria)
                if not hit and cell_value is _TEXT_VALUE and criteria[0] == "text":
                    # 文本 criteria：经 index 取该格原始文本做大小写不敏感匹配。
                    if index is not None:
                        info = index.cell(f"{rs}!{ra}")
                        if info is not None and not info.is_formula \
                                and isinstance(info.value, str):
                            hit = info.value.casefold() == criteria[1].casefold()
                if hit:
                    matched += 1
                    value = values.get(f"{vs}!{va}", _BLANK)
                    if isinstance(value, (int, float)):
                        total += float(value)
            if name == "COUNTIF":
                return float(matched)
            return total
        elif name == "VLOOKUP":
            if len(ast[2]) != 4:
                raise ValueError(
                    "VLOOKUP requires an explicit FALSE as the 4th argument"
                )
            if ast[2][3][0] != "bool" or ast[2][3][1] is not False:
                raise ValueError(
                    "VLOOKUP only supports exact match (4th argument FALSE)"
                )
            lookup = _lookup_target_value(ast[2][0], values, index)
            table_arg = ast[2][1]
            if table_arg[0] != "cellref" or ":" not in table_arg[1]:
                raise ValueError("VLOOKUP table must be a range reference")
            col_index = int(_numeric_of(_eval_ast(ast[2][2], values, index)))
            if col_index < 1:
                raise ValueError(f"VLOOKUP col_index must be >= 1, got {col_index}")
            sheet, _, rest = table_arg[1].partition("!")
            clean = _sanitize_range_str(rest)
            try:
                min_col, min_row, max_col, max_row = range_boundaries(clean)
            except Exception as exc:
                raise ValueError(f"VLOOKUP table range invalid: {table_arg[1]!r}") from exc
            if col_index > max_col - min_col + 1:
                raise ValueError(
                    f"VLOOKUP col_index {col_index} exceeds table width"
                )
            for row in range(min_row, max_row + 1):
                key_addr = f"{sheet}!{get_column_letter(min_col)}{row}"
                if _lookup_equal(
                    lookup, _cell_lookup_value(key_addr, values, index)
                ):
                    result_addr = (
                        f"{sheet}!"
                        f"{get_column_letter(min_col + col_index - 1)}{row}"
                    )
                    raw = values.get(result_addr, _BLANK)
                    if isinstance(raw, (int, float)):
                        return float(raw)
                    if raw is _TEXT_VALUE:
                        if index is not None:
                            info = index.cell(result_addr)
                            if (
                                info is not None and not info.is_formula
                                and isinstance(info.value, str)
                            ):
                                return _TEXT_VALUE
                        return _TEXT_VALUE
                    return 0.0  # Excel：命中空白格 → 0
            raise ValueError(
                f"VLOOKUP lookup value not found (#N/A): {lookup!r}"
            )
        else:
            raise ValueError(f"Unsupported function: {name}")
    else:
        raise ValueError(f"Unknown AST node: {ast}")


# Marker for text constants. They are ignored inside SUM ranges, but a direct
# arithmetic reference must fail rather than silently becoming zero.
_TEXT_VALUE = object()
# Marker for blank cells (empty or absent). Aggregations skip them; plain
# arithmetic treats them as 0, matching Excel's blank-cell semantics.
_BLANK = object()


def _numeric_of(value: float | object) -> float:
    """把求值结果数值化：文本参与数值语境直接失败，空白按 0。

    示例：IF 的条件、比较运算的操作数都经过它——与 Excel 的
    文本不参与算术语义一致，避免静默当 0。
    """
    if value is _TEXT_VALUE:
        raise ValueError("text value used in a numeric context")
    if value is _BLANK:
        return 0.0
    return float(value)


def _parse_criteria(
    arg, values: dict[str, float | object], index: WorkbookIndex | None
) -> tuple[str, float | str] | tuple[str, str, float]:
    """解析 SUMIF/COUNTIF 的 criteria 实参为比较目标。

    形态：("number", v) 相等匹配 / ("text", s) 文本精确（casefold）/
    ("op", op, n) 数值比较（op ∈ > < >= <= = <>）。
    通配符 * ? 显式拒绝；非法形态 raise。
    """
    if arg[0] == "string":
        text = arg[1]
        if "*" in text or "?" in text:
            raise ValueError(
                f"wildcards are not supported in criteria: {text!r}"
            )
        # 先长符号（>=、<=、<>）再短符号（>、<、=），避免 ">=" 被截成 ">"。
        # 数值后缀转换失败时 break，回落到文本精确匹配（如 ">abc"）。
        for op_str in (">=", "<=", "<>", ">", "<", "="):
            if text.startswith(op_str):
                try:
                    return ("op", op_str, float(text[len(op_str):]))
                except ValueError:
                    break
        if text == "":
            return ("number", 0.0)
        try:
            return ("number", float(text))
        except ValueError:
            return ("text", text)
    value = _eval_ast(arg, values, index)
    value = _numeric_of(value)
    return ("number", value)


def _criteria_match(cell_value: float | object, criteria) -> bool:
    """单个格是否命中 criteria（空白格与文本格不匹配任何数值 criteria）。

    数值格按 criteria 形态判定；文本 criteria 用 casefold 大小写不敏感
    比较，支持 values 里直接存放的字符串（纯求值路径，evaluate_all 会把
    字符串常量转成 _TEXT_VALUE，那条路径经 index 回取原文）。
    """
    if cell_value is _BLANK:
        return False
    kind = criteria[0]
    if kind == "text":
        if isinstance(cell_value, str):
            return cell_value.casefold() == criteria[1].casefold()
        return False
    if not isinstance(cell_value, (int, float)):
        return False  # 文本格（_TEXT_VALUE 或 str）不参与数值比较
    number = float(cell_value)
    if kind == "number":
        return number == criteria[1]
    op = criteria[1]
    if op == ">":
        return number > criteria[2]
    if op == "<":
        return number < criteria[2]
    if op == ">=":
        return number >= criteria[2]
    if op == "<=":
        return number <= criteria[2]
    if op == "<>":
        return number != criteria[2]
    return number == criteria[2]  # "="


def _lookup_target_value(
    arg, values: dict[str, float | object], index: WorkbookIndex | None
) -> float | str:
    """提取 VLOOKUP 的查找值：字符串字面量原样；数值求值；文本格经 index 取原文。"""
    if arg[0] == "string":
        return arg[1]
    value = _eval_ast(arg, values, index)
    if value is _TEXT_VALUE:
        if arg[0] == "cellref" and index is not None:
            info = index.cell(arg[1])
            if info is not None and not info.is_formula and isinstance(info.value, str):
                return info.value
        raise ValueError("VLOOKUP lookup value must be numeric or a text cell")
    return _numeric_of(value)


def _cell_lookup_value(
    addr: str, values: dict[str, float | object], index: WorkbookIndex | None
) -> float | str | None:
    """取查找列某格的原始比较值：数字 → float；文本常量 → str；空白/文本结果 → None。"""
    raw = values.get(addr, _BLANK)
    if raw is _BLANK:
        return None
    if raw is _TEXT_VALUE:
        if index is not None:
            info = index.cell(addr)
            if info is not None and not info.is_formula and isinstance(info.value, str):
                return info.value
        return None  # 公式格的文本结果无法回取原文 → 视为不可比较
    if isinstance(raw, str):
        # 直调 _eval_ast 路径：values 里可存真实字符串（与 Task 5 一致）。
        return raw
    return float(raw)


def _lookup_equal(lookup: float | str, target: float | str | None) -> bool:
    """精确相等：数字比数字；文本 casefold 相等（Excel VLOOKUP 不区分大小写）。"""
    if isinstance(lookup, str):
        return isinstance(target, str) and lookup.casefold() == target.casefold()
    return isinstance(target, float) and float(lookup) == target


def _range_cells(full_addr: str) -> list[tuple[str, str]]:
    """范围地址 → 行主序 (sheet, addr) 列表（SUMIF/VLOOKUP 共用）。"""
    sheet, _, rest = full_addr.partition("!")
    return expand_range(sheet, rest)


def _cell_or_range_numerics(
    full_addr: str, values: dict[str, float | object], index: WorkbookIndex | None = None
) -> list[float]:
    """按 Excel AVERAGE/COUNT 语义取出"参与计算"的数字。

    - 范围引用逐格展开，单格引用原样处理；
    - 借助 index 区分空白格（未写入或值为 None → 跳过）；index 不可用时
      退化为 values 字典语义（缺失格按 0.0 计入）；
    - 公式格取重算值；文本常量（_TEXT_VALUE）与空白（_BLANK）跳过。
    """
    if ":" in full_addr:
        addrs = [f"{s}!{a}" for s, a in _range_cells(full_addr)]
    else:
        addrs = [full_addr]
    numbers: list[float] = []
    for addr in addrs:
        if index is not None:
            info = index.cell(addr)
            if info is None or (not info.is_formula and info.value is None):
                continue  # 空白格不计入（Excel AVERAGE/COUNT 语义）
        value = values.get(addr, 0.0)
        if value is _TEXT_VALUE or value is _BLANK:
            continue
        numbers.append(float(value))
    return numbers


class RecalcEngine:
    """受限 Excel 公式重算引擎。

    工作流程：
        1. 从 WorkbookIndex 中读取公式和单元格数据；
        2. 使用 DependencyGraph 确定公式的计算顺序；
        3. 使用 _ExprParser 将公式解析为 AST；
        4. 使用 _eval_ast() 计算 AST，得到公式结果。

    示例：先构建 `graph`，再创建 `RecalcEngine(index, graph)`，
    最后调用 `evaluate_all()` 或 `evaluate_cell()` 进行重算。
    """

    def __init__(self, index: WorkbookIndex, graph: DependencyGraph):
        """初始化重算引擎。

        `index` 提供单元格数据，`graph` 提供公式依赖关系。
        """
        self.index = index
        self.graph = graph
        # 最近一次 evaluate_all/evaluate_cell 中每个失败公式的根因
        # （"异常类型: 消息"）。验证器用它把"无法重算"变成可行动的
        # 反馈（如 Unsupported function: ROUND）；成功清空。
        self.errors: dict[str, str] = {}

    @staticmethod
    def _constant_value(value: object) -> float | object:
        """转换常量值，并保留文本/空白/数字之间的区别。

        示例：`100` 转为 `100.0`；`None`（空白格）转为 `_BLANK` 标记，
        供聚合函数跳过、算术按 0 处理；`"文本"` 转为文本标记。
        这样 SUM 可以忽略文本，而 `="文本"+1` 不会被错误地当成 `0+1`。
        """
        if value is None:
            return _BLANK
        if isinstance(value, (int, float, bool)):
            return float(value)
        return _TEXT_VALUE

    def evaluate_all(self) -> dict[str, float]:
        """按照拓扑顺序重新计算整个工作簿。

        示例：若 `A1=10`、`B1=20`、`C1=A1+B1`，返回结果中会包含
        `{"Sheet1!C1": 30.0}`。计算失败的公式记为 `NaN`，并在
        ``self.errors`` 里记录根因（异常类型 + 消息）。

        multi-cell 兼容：工作簿存在既有循环依赖时，`topological_order()`
        为空、所有公式都无法求值；这里退化为按 SCC 缩点图计算无循环
        部分，循环成员记为 `NaN`，让验证器能把既有循环与新引入的
        循环区分开。
        """
        self.errors = {}
        values: dict[str, float] = {}
        for sheet in self.index.sheets:
            for addr, info in sheet.cells.items():
                if not info.is_formula:
                    values[info.ref.full_address] = self._constant_value(info.value)
        order, cyclic = self.graph.evaluation_order()
        for cell_addr in order:
            info = self.index.cell(cell_addr)
            if info and info.is_formula:
                try:
                    values[cell_addr] = self._eval_formula(
                        info.formula, info.ref, values
                    )
                except Exception as exc:
                    self.errors[cell_addr] = f"{type(exc).__name__}: {exc}"[:200]
                    values[cell_addr] = float("nan")
        for cell_addr in cyclic:
            info = self.index.cell(cell_addr)
            if info and info.is_formula:
                self.errors[cell_addr] = "cyclic dependency"
                values[cell_addr] = float("nan")
        return values

    def evaluate_cell(self, cell_addr: str) -> float:
        """计算指定单元格，并先计算它的所有上游依赖。

        示例：调用 `evaluate_cell("Sheet1!D1")` 时，会先计算 D1 依赖的
        C1、B1 等公式，最后返回 D1 的数值。
        """
        order = self.graph.topological_order()
        self.errors = {}
        values: dict[str, float] = {}
        for sheet in self.index.sheets:
            for addr, info in sheet.cells.items():
                if not info.is_formula:
                    values[info.ref.full_address] = self._constant_value(info.value)
        for cell in order:
            info = self.index.cell(cell)
            if info and info.is_formula:
                try:
                    values[cell] = self._eval_formula(
                        info.formula, info.ref, values
                    )
                except Exception as exc:
                    self.errors[cell] = f"{type(exc).__name__}: {exc}"[:200]
                    values[cell] = float("nan")
            if cell == cell_addr:
                break
        return values.get(cell_addr, float("nan"))

    def _eval_formula(
        self, formula: str, host: CellRef, values: dict[str, float]
    ) -> float:
        """解析并计算一个公式。

        示例：`_eval_formula("=A1+B1", host, values)` 会先分词、生成 AST，
        再根据 `values` 中的 A1/B1 值返回计算结果。
        """
        tokens = _tokenize_openpyxl(formula, host)
        parser = _ExprParser(tokens)
        ast = parser.parse()
        return _eval_ast(ast, values, self.index)

    def evaluate_expression(self, formula: str, host: CellRef) -> float:
        """在工作簿当前值上求值一个表达式，不写入任何单元格。

        供值反推提示等调用方复用引擎的求值口径；公式无法求值时异常
        向上抛出（调用方自行决定如何处理）。
        """
        return self._eval_formula(formula, host, self.evaluate_all())